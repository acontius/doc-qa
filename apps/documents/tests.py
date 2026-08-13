import math
import os
import shutil
import tempfile
from contextlib import contextmanager
from io import BytesIO
from pathlib import Path
from unittest.mock import Mock, patch

from rest_framework.test import (
    APIClient,
    APITestCase,
)
import numpy as np

from django.conf import settings
from django.core.files.storage import default_storage
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import connection
from django.test import (
    TestCase,
    override_settings,
)
from docx import Document as DocxDocument
from rest_framework import status

from .models import Answer, Document, DocumentChunk, Question
from .services import (
    DocumentIngestionError,
    DocumentRetrievalError,
    EmbeddingConfigurationError,
    EmbeddingService,
    compute_keyword_score,
    compute_vector_score,
    create_document,
    delete_document,
    extract_docx_text,
    get_embedding_service,
    reindex_document,
    retrieve_chunks,
    split_document_text,
    tokenize_text,
    update_document,
    validate_docx_upload,
    answer_question,
    ask_and_persist_question,
    build_rag_context,
    get_chat_model,
    is_llm_configured,
    parse_llm_response,
    NO_CONTEXT_ANSWER,
    RAGError,
    RetrievedChunk,
)


class DocumentTestMixin:
    def make_docx(
        self,
        paragraphs: list[str],
        table_rows: list[list[str]] | None = None,
    ) -> bytes:
        document = DocxDocument()

        for paragraph in paragraphs:
            document.add_paragraph(paragraph)

        if table_rows:
            table = document.add_table(
                rows=1,
                cols=len(table_rows[0]),
            )

            for row_index, row in enumerate(table_rows):
                if row_index == 0:
                    cells = table.rows[0].cells
                else:
                    cells = table.add_row().cells

                for cell, value in zip(
                    cells,
                    row,
                    strict=True,
                ):
                    cell.text = value

        buffer = BytesIO()
        document.save(buffer)

        return buffer.getvalue()

    def make_upload(
        self,
        content: bytes,
        name: str = "sample.docx",
    ) -> SimpleUploadedFile:
        return SimpleUploadedFile(
            name,
            content,
            content_type=(
                "application/"
                "vnd.openxmlformats-officedocument."
                "wordprocessingml.document"
            ),
        )

    @contextmanager
    def mocked_embedding_service(
        self,
        model_name: str = "test-model",
        value: float = 0.01,
    ):
        embedding_service = Mock()
        embedding_service.model_name = model_name

        embedding_service.embed_documents.side_effect = (
            lambda texts: [
                [value] * settings.EMBEDDING_DIMENSION
                for _ in texts
            ]
        )

        with patch(
            "apps.documents.services.get_embedding_service",
            return_value=embedding_service,
        ):
            yield embedding_service


class TemporaryMediaMixin:
    @classmethod
    def setUpClass(cls):
        cls._media_root = Path(
            tempfile.mkdtemp(
                prefix="doc_qa_test_media_",
            )
        )

        cls._media_override = override_settings(
            MEDIA_ROOT=str(cls._media_root),
        )
        cls._media_override.enable()

        # Django caches storage instances and LazyObject wrappers.
        # Reset both so the new MEDIA_ROOT is used for file I/O.
        from django.core.files.storage import (
            default_storage,
            storages,
        )
        from django.utils.functional import empty

        storages._storages.clear()
        default_storage._wrapped = empty

        super().setUpClass()

    @classmethod
    def tearDownClass(cls):
        try:
            super().tearDownClass()
        finally:
            cls._media_override.disable()
            from django.core.files.storage import (
                default_storage,
                storages,
            )
            from django.utils.functional import empty

            storages._storages.clear()
            default_storage._wrapped = empty
            shutil.rmtree(
                cls._media_root,
                ignore_errors=True,
            )


class DocumentValidationTests(
    DocumentTestMixin,
    TestCase,
):
    def test_rejects_non_docx_extension(self):
        uploaded_file = self.make_upload(
            b"not a docx",
            name="sample.pdf",
        )

        with self.assertRaises(DocumentIngestionError) as context:
            validate_docx_upload(uploaded_file)

        self.assertEqual(
            context.exception.code,
            "invalid_extension",
        )

    def test_rejects_empty_file(self):
        uploaded_file = self.make_upload(
            b"",
        )

        with self.assertRaises(DocumentIngestionError) as context:
            validate_docx_upload(uploaded_file)

        self.assertEqual(
            context.exception.code,
            "empty_file",
        )

    def test_rejects_oversized_file(self):
        content = b"x" * (
            settings.MAX_UPLOAD_SIZE_MB * 1024 * 1024 + 1
        )

        uploaded_file = self.make_upload(
            content,
        )

        with self.assertRaises(DocumentIngestionError) as context:
            validate_docx_upload(uploaded_file)

        self.assertEqual(
            context.exception.code,
            "file_too_large",
        )

    def test_rejects_invalid_docx_structure(self):
        buffer = BytesIO()

        import zipfile

        with zipfile.ZipFile(
            buffer,
            "w",
            compression=zipfile.ZIP_DEFLATED,
        ) as archive:
            archive.writestr(
                "[Content_Types].xml",
                "<Types></Types>",
            )

        uploaded_file = self.make_upload(
            buffer.getvalue(),
        )

        with self.assertRaises(DocumentIngestionError) as context:
            validate_docx_upload(uploaded_file)

        self.assertEqual(
            context.exception.code,
            "invalid_docx_structure",
        )

    def test_rejects_corrupted_docx(self):
        uploaded_file = self.make_upload(
            b"this is not a zip file",
        )

        with self.assertRaises(DocumentIngestionError) as context:
            validate_docx_upload(uploaded_file)

        self.assertEqual(
            context.exception.code,
            "invalid_docx",
        )

    def test_validation_preserves_file_position(self):
        content = self.make_docx(
            paragraphs=["متن آزمایشی"],
        )

        uploaded_file = self.make_upload(content)

        uploaded_file.seek(5)

        validate_docx_upload(uploaded_file)

        self.assertEqual(
            uploaded_file.tell(),
            0,
        )


class DocumentExtractionTests(
    DocumentTestMixin,
    TestCase,
):
    def test_extracts_paragraphs_and_tables(self):
        content = self.make_docx(
            paragraphs=[
                "اولین پاراگراف",
                "پاراگراف دوم",
            ],
            table_rows=[
                ["نام", "مقدار"],
                ["A", "100"],
            ],
        )

        uploaded_file = self.make_upload(content)

        text = extract_docx_text(
            uploaded_file,
        )

        self.assertIn(
            "اولین پاراگراف",
            text,
        )

        self.assertIn(
            "پاراگراف دوم",
            text,
        )

        self.assertIn(
            "نام | مقدار",
            text,
        )

        self.assertIn(
            "A | 100",
            text,
        )

    def test_rejects_empty_document(self):
        document = DocxDocument()

        buffer = BytesIO()
        document.save(buffer)

        uploaded_file = self.make_upload(
            buffer.getvalue(),
        )

        with self.assertRaises(DocumentIngestionError) as context:
            extract_docx_text(uploaded_file)

        self.assertEqual(
            context.exception.code,
            "empty_document",
        )

    def test_normalizes_document_text(self):
        content = self.make_docx(
            paragraphs=[
                "  متن    با فاصله‌های زیاد  ",
            ],
        )

        uploaded_file = self.make_upload(content)

        text = extract_docx_text(
            uploaded_file,
        )

        self.assertEqual(
            text,
            "متن با فاصله‌های زیاد",
        )


class ChunkingTests(TestCase):
    def test_chunking_removes_empty_chunks(self):
        text = (
            "این یک متن آزمایشی است. " * 300
        )

        chunks = split_document_text(text)

        self.assertTrue(chunks)

        self.assertTrue(
            all(
                chunk.strip()
                for chunk in chunks
            )
        )

    def test_chunking_respects_configured_size(self):
        text = "A" * 3500

        chunks = split_document_text(text)

        self.assertTrue(chunks)

        self.assertTrue(
            all(
                len(chunk) <= settings.CHUNK_SIZE
                for chunk in chunks
            )
        )

    def test_chunking_uses_overlap(self):
        text = "A" * 2500

        chunks = split_document_text(text)

        self.assertGreaterEqual(
            len(chunks),
            2,
        )

        self.assertEqual(
            chunks[0][-settings.CHUNK_OVERLAP:],
            chunks[1][:settings.CHUNK_OVERLAP],
        )

    @override_settings(
        CHUNK_SIZE=0,
    )
    def test_rejects_invalid_chunk_size(self):
        with self.assertRaises(DocumentIngestionError) as context:
            split_document_text("some text")

        self.assertEqual(
            context.exception.code,
            "invalid_chunk_size",
        )

    @override_settings(
        CHUNK_SIZE=1000,
        CHUNK_OVERLAP=1000,
    )
    def test_rejects_invalid_chunk_overlap(self):
        with self.assertRaises(DocumentIngestionError) as context:
            split_document_text("some text")

        self.assertEqual(
            context.exception.code,
            "invalid_chunk_overlap",
        )


class EmbeddingServiceTests(TestCase):
    def tearDown(self):
        get_embedding_service.cache_clear()
        super().tearDown()

    @patch("apps.documents.services.SentenceTransformer")
    def test_loads_configured_model(
        self,
        mock_sentence_transformer,
    ):
        model = Mock()

        model.get_embedding_dimension.return_value = (
            settings.EMBEDDING_DIMENSION
        )

        mock_sentence_transformer.return_value = model

        service = EmbeddingService()

        mock_sentence_transformer.assert_called_once_with(
            settings.EMBEDDING_MODEL,
        )

        self.assertEqual(
            service.model_name,
            settings.EMBEDDING_MODEL,
        )

    @patch("apps.documents.services.SentenceTransformer")
    def test_rejects_embedding_dimension_mismatch(
        self,
        mock_sentence_transformer,
    ):
        model = Mock()

        model.get_embedding_dimension.return_value = (
            settings.EMBEDDING_DIMENSION + 1
        )

        mock_sentence_transformer.return_value = model

        with self.assertRaises(
            EmbeddingConfigurationError,
        ) as context:
            EmbeddingService()

        self.assertEqual(
            context.exception.code,
            "embedding_dimension_mismatch",
        )

    @patch("apps.documents.services.SentenceTransformer")
    def test_handles_model_load_failure(
        self,
        mock_sentence_transformer,
    ):
        mock_sentence_transformer.side_effect = RuntimeError(
            "model unavailable",
        )

        with self.assertRaises(
            EmbeddingConfigurationError,
        ) as context:
            EmbeddingService()

        self.assertEqual(
            context.exception.code,
            "embedding_model_load_failed",
        )

    @patch("apps.documents.services.SentenceTransformer")
    def test_embeds_multiple_documents(
        self,
        mock_sentence_transformer,
    ):
        model = Mock()

        model.get_embedding_dimension.return_value = (
            settings.EMBEDDING_DIMENSION
        )

        vectors = np.zeros(
            (
                2,
                settings.EMBEDDING_DIMENSION,
            ),
            dtype=np.float32,
        )

        vectors[0][0] = 1.0
        vectors[1][1] = 1.0

        model.encode.return_value = vectors

        mock_sentence_transformer.return_value = model

        service = EmbeddingService()

        result = service.embed_documents(
            [
                "متن اول",
                "متن دوم",
            ]
        )

        self.assertEqual(
            len(result),
            2,
        )

        self.assertEqual(
            len(result[0]),
            settings.EMBEDDING_DIMENSION,
        )

        self.assertEqual(
            len(result[1]),
            settings.EMBEDDING_DIMENSION,
        )

        model.encode.assert_called_once()

        encode_kwargs = (
            model.encode.call_args.kwargs
        )

        self.assertTrue(
            encode_kwargs["normalize_embeddings"],
        )

        self.assertTrue(
            encode_kwargs["convert_to_numpy"],
        )

    @patch("apps.documents.services.SentenceTransformer")
    def test_rejects_invalid_embedding_shape(
        self,
        mock_sentence_transformer,
    ):
        model = Mock()

        model.get_embedding_dimension.return_value = (
            settings.EMBEDDING_DIMENSION
        )

        model.encode.return_value = np.zeros(
            settings.EMBEDDING_DIMENSION,
            dtype=np.float32,
        )

        mock_sentence_transformer.return_value = model

        service = EmbeddingService()

        with self.assertRaises(
            EmbeddingConfigurationError,
        ) as context:
            service.embed_documents(
                ["text"],
            )

        self.assertEqual(
            context.exception.code,
            "invalid_embedding_shape",
        )

    @patch("apps.documents.services.SentenceTransformer")
    def test_empty_input_returns_empty_output(
        self,
        mock_sentence_transformer,
    ):
        model = Mock()

        model.get_embedding_dimension.return_value = (
            settings.EMBEDDING_DIMENSION
        )

        mock_sentence_transformer.return_value = model

        service = EmbeddingService()

        result = service.embed_documents([])

        self.assertEqual(
            result,
            [],
        )

        model.encode.assert_not_called()

    @patch("apps.documents.services.EmbeddingService")
    def test_embedding_service_is_singleton(
        self,
        mock_embedding_service,
    ):
        get_embedding_service.cache_clear()

        first = Mock()
        second = first

        mock_embedding_service.return_value = first

        service_one = get_embedding_service()
        service_two = get_embedding_service()

        self.assertIs(
            service_one,
            service_two,
        )

        mock_embedding_service.assert_called_once()


class DocumentServiceTests(
    DocumentTestMixin,
    TemporaryMediaMixin,
    TestCase,
):
    def test_create_document_creates_complete_state(self):
        with self.mocked_embedding_service():
            content = self.make_docx(
                paragraphs=[
                    "متن اول " * 100,
                    "متن دوم " * 100,
                ],
            )

            document = create_document(
                title="Test Document",
                uploaded_file=self.make_upload(
                    content,
                ),
            )

        self.assertEqual(
            document.title,
            "Test Document",
        )

        self.assertTrue(
            document.full_text,
        )

        self.assertGreater(
            document.chunk_count,
            0,
        )

        self.assertEqual(
            document.chunks.count(),
            document.chunk_count,
        )

        self.assertTrue(
            document.file.name.startswith(
                "documents/",
            )
        )

        self.assertTrue(
            document.file.name.endswith(
                ".docx",
            )
        )

        first_chunk = (
            document.chunks.order_by(
                "chunk_index",
            ).first()
        )

        self.assertIsNotNone(
            first_chunk,
        )

        self.assertEqual(
            len(first_chunk.embedding),
            settings.EMBEDDING_DIMENSION,
        )

        self.assertEqual(
            document.embedding_model,
            settings.EMBEDDING_MODEL,
        )

        self.assertTrue(
            default_storage.exists(
                document.file.name,
            )
        )

    def test_create_document_stores_vector_in_postgres(self):
        with self.mocked_embedding_service():
            content = self.make_docx(
                paragraphs=[
                    "متن برای تست pgvector " * 150,
                ],
            )

            document = create_document(
                title="Vector Test",
                uploaded_file=self.make_upload(
                    content,
                ),
            )

        chunk = (
            document.chunks
            .order_by("chunk_index")
            .first()
        )

        with connection.cursor() as cursor:
            cursor.execute(
                """
                SELECT vector_dims(embedding)
                FROM documents_documentchunk
                WHERE id = %s
                """,
                [chunk.id],
            )

            dimension = cursor.fetchone()[0]

        self.assertEqual(
            dimension,
            settings.EMBEDDING_DIMENSION,
        )

    def test_create_document_rolls_back_database_and_file(
        self,
    ):
        # Other tests in this class may leave files on disk
        # (TestCase rolls back DB rows, not filesystem writes).
        files_before = {
            path
            for path in Path(settings.MEDIA_ROOT).rglob("*")
            if path.is_file()
        }

        with self.mocked_embedding_service():
            content = self.make_docx(
                paragraphs=[
                    "متن rollback " * 100,
                ],
            )

            with patch(
                "apps.documents.services._replace_chunks",
                side_effect=RuntimeError(
                    "forced persistence failure",
                ),
            ):
                with self.assertRaises(
                    RuntimeError,
                ):
                    create_document(
                        title="Rollback Test",
                        uploaded_file=self.make_upload(
                            content,
                        ),
                    )

        self.assertEqual(
            Document.objects.count(),
            0,
        )

        self.assertEqual(
            DocumentChunk.objects.count(),
            0,
        )

        files_after = {
            path
            for path in Path(settings.MEDIA_ROOT).rglob("*")
            if path.is_file()
        }

        # The failed create must not leave a new orphaned file.
        self.assertEqual(
            files_after,
            files_before,
        )

    def test_update_title_does_not_reindex(self):
        with self.mocked_embedding_service() as embedding_service:
            content = self.make_docx(
                paragraphs=[
                    "متن اصلی " * 100,
                ],
            )

            document = create_document(
                title="Original",
                uploaded_file=self.make_upload(
                    content,
                ),
            )

            embedding_service.embed_documents.reset_mock()

            updated = update_document(
                document,
                title="Updated",
            )

        self.assertEqual(
            updated.title,
            "Updated",
        )

        self.assertEqual(
            embedding_service.embed_documents.call_count,
            0,
        )

        self.assertGreater(
            updated.chunks.count(),
            0,
        )

    def test_document_replacement_removes_old_chunks_and_file(
        self,
    ):
        with self.mocked_embedding_service():
            old_content = self.make_docx(
                paragraphs=[
                    "متن اولیه " * 100,
                ],
            )

            document = create_document(
                title="Original",
                uploaded_file=self.make_upload(
                    old_content,
                ),
            )

            old_file_name = document.file.name
            old_chunk_ids = set(
                document.chunks.values_list(
                    "id",
                    flat=True,
                )
            )

            new_content = self.make_docx(
                paragraphs=[
                    "متن جدید " * 200,
                ],
            )

            updated = update_document(
                document,
                title="Updated",
                uploaded_file=self.make_upload(
                    new_content,
                ),
            )

        new_chunk_ids = set(
            updated.chunks.values_list(
                "id",
                flat=True,
            )
        )

        self.assertEqual(
            updated.title,
            "Updated",
        )

        self.assertTrue(
            updated.full_text.startswith(
                "متن جدید",
            )
        )

        self.assertTrue(
            new_chunk_ids,
        )

        self.assertTrue(
            old_chunk_ids.isdisjoint(
                new_chunk_ids,
            )
        )

        self.assertEqual(
            updated.chunk_count,
            updated.chunks.count(),
        )

        self.assertNotEqual(
            old_file_name,
            updated.file.name,
        )

        self.assertFalse(
            default_storage.exists(
                old_file_name,
            )
        )

        self.assertTrue(
            default_storage.exists(
                updated.file.name,
            )
        )

    def test_reindex_rebuilds_chunks_from_existing_file(
        self,
    ):
        with self.mocked_embedding_service():
            content = self.make_docx(
                paragraphs=[
                    "متن قابل reindex " * 200,
                ],
            )

            document = create_document(
                title="Reindex Test",
                uploaded_file=self.make_upload(
                    content,
                ),
            )

            original_file_name = document.file.name
            original_chunk_ids = set(
                document.chunks.values_list(
                    "id",
                    flat=True,
                )
            )

            reindexed = reindex_document(
                document,
            )

        new_chunk_ids = set(
            reindexed.chunks.values_list(
                "id",
                flat=True,
            )
        )

        self.assertEqual(
            reindexed.file.name,
            original_file_name,
        )

        self.assertTrue(
            original_chunk_ids.isdisjoint(
                new_chunk_ids,
            )
        )

        self.assertEqual(
            reindexed.chunk_count,
            reindexed.chunks.count(),
        )

        self.assertTrue(
            reindexed.full_text,
        )

        self.assertTrue(
            default_storage.exists(
                reindexed.file.name,
            )
        )

    def test_reindex_reports_missing_source_file(
        self,
    ):
        with self.mocked_embedding_service():
            content = self.make_docx(
                paragraphs=[
                    "متن فایل " * 100,
                ],
            )

            document = create_document(
                title="Missing Source",
                uploaded_file=self.make_upload(
                    content,
                ),
            )

            old_chunk_ids = set(
                document.chunks.values_list(
                    "id",
                    flat=True,
                )
            )

            default_storage.delete(
                document.file.name,
            )

            with self.assertRaises(
                DocumentIngestionError,
            ) as context:
                reindex_document(document)

        self.assertEqual(
            context.exception.code,
            "source_file_missing",
        )

        self.assertEqual(
            set(
                DocumentChunk.objects.filter(
                    document=document,
                ).values_list(
                    "id",
                    flat=True,
                )
            ),
            old_chunk_ids,
        )

    def test_delete_document_removes_database_rows_and_file(
        self,
    ):
        with self.mocked_embedding_service():
            content = self.make_docx(
                paragraphs=[
                    "متن برای حذف " * 100,
                ],
            )

            document = create_document(
                title="Delete Test",
                uploaded_file=self.make_upload(
                    content,
                ),
            )

        file_name = document.file.name

        self.assertTrue(
            default_storage.exists(
                file_name,
            )
        )

        delete_document(document)

        self.assertFalse(
            Document.objects.filter(
                pk=document.pk,
            ).exists()
        )

        self.assertFalse(
            DocumentChunk.objects.filter(
                document_id=document.pk,
            ).exists()
        )

        self.assertFalse(
            default_storage.exists(
                file_name,
            )
        )


class DocumentAPITests(
    DocumentTestMixin,
    TemporaryMediaMixin,
    APITestCase,
):
    def setUp(self):
        self.client = APIClient()

    def test_create_document_api(self):
        with self.mocked_embedding_service():
            content = self.make_docx(
                paragraphs=[
                    "محتوای سند آزمایشی " * 100,
                ],
            )

            response = self.client.post(
                "/api/documents/",
                {
                    "title": "سند آزمایشی",
                    "file": self.make_upload(content),
                },
                format="multipart",
            )

        self.assertEqual(
            response.status_code,
            status.HTTP_201_CREATED,
        )

        document = Document.objects.get(
            pk=response.data["id"],
        )

        self.assertEqual(
            document.chunk_count,
            document.chunks.count(),
        )

        self.assertEqual(
            response.data["chunk_count"],
            document.chunk_count,
        )

    def test_create_requires_file(self):
        response = self.client.post(
            "/api/documents/",
            {
                "title": "بدون فایل",
            },
            format="multipart",
        )

        self.assertEqual(
            response.status_code,
            status.HTTP_400_BAD_REQUEST,
        )

        self.assertEqual(
            response.data["error"]["code"],
            "validation_error",
        )

    def test_reject_invalid_file(self):
        response = self.client.post(
            "/api/documents/",
            {
                "title": "Invalid",
                "file": self.make_upload(
                    b"not a docx",
                    name="evil.exe",
                ),
            },
            format="multipart",
        )

        self.assertEqual(
            response.status_code,
            status.HTTP_400_BAD_REQUEST,
        )

        self.assertEqual(
            response.data["error"]["code"],
            "validation_error",
        )

    def test_list_documents(self):
        with self.mocked_embedding_service():
            content = self.make_docx(
                paragraphs=[
                    "hello " * 100,
                ],
            )

            self.client.post(
                "/api/documents/",
                {
                    "title": "Document One",
                    "file": self.make_upload(content),
                },
                format="multipart",
            )

        response = self.client.get(
            "/api/documents/",
        )

        self.assertEqual(
            response.status_code,
            status.HTTP_200_OK,
        )

        self.assertEqual(
            len(response.data["results"]),
            1,
        )

    def test_retrieve_document(self):
        with self.mocked_embedding_service():
            content = self.make_docx(
                paragraphs=[
                    "متن برای retrieve " * 100,
                ],
            )

            create_response = self.client.post(
                "/api/documents/",
                {
                    "title": "Retrieve Test",
                    "file": self.make_upload(content),
                },
                format="multipart",
            )

        document_id = create_response.data["id"]

        response = self.client.get(
            f"/api/documents/{document_id}/",
        )

        self.assertEqual(
            response.status_code,
            status.HTTP_200_OK,
        )

        self.assertEqual(
            response.data["id"],
            document_id,
        )

        self.assertIn(
            "متن برای retrieve",
            response.data["full_text"],
        )

    def test_patch_title(self):
        with self.mocked_embedding_service():
            content = self.make_docx(
                paragraphs=[
                    "متن patch " * 100,
                ],
            )

            create_response = self.client.post(
                "/api/documents/",
                {
                    "title": "Original",
                    "file": self.make_upload(content),
                },
                format="multipart",
            )

        document_id = create_response.data["id"]

        response = self.client.patch(
            f"/api/documents/{document_id}/",
            {
                "title": "Updated",
            },
            format="multipart",
        )

        self.assertEqual(
            response.status_code,
            status.HTTP_200_OK,
        )

        self.assertEqual(
            response.data["title"],
            "Updated",
        )

    def test_patch_file_reindexes_document(self):
        with self.mocked_embedding_service():
            old_content = self.make_docx(
                paragraphs=[
                    "متن قدیمی " * 100,
                ],
            )

            create_response = self.client.post(
                "/api/documents/",
                {
                    "title": "Original",
                    "file": self.make_upload(
                        old_content,
                    ),
                },
                format="multipart",
            )

            document_id = create_response.data["id"]

            new_content = self.make_docx(
                paragraphs=[
                    "متن جدید " * 150,
                ],
            )

            response = self.client.patch(
                f"/api/documents/{document_id}/",
                {
                    "title": "Updated",
                    "file": self.make_upload(
                        new_content,
                    ),
                },
                format="multipart",
            )

        self.assertEqual(
            response.status_code,
            status.HTTP_200_OK,
        )

        self.assertEqual(
            response.data["title"],
            "Updated",
        )

        self.assertIn(
            "متن جدید",
            response.data["full_text"],
        )

        self.assertNotIn(
            "متن قدیمی",
            response.data["full_text"],
        )

    def test_reindex_endpoint(self):
        with self.mocked_embedding_service():
            content = self.make_docx(
                paragraphs=[
                    "متن reindex endpoint " * 100,
                ],
            )

            create_response = self.client.post(
                "/api/documents/",
                {
                    "title": "Reindex API",
                    "file": self.make_upload(content),
                },
                format="multipart",
            )

            document_id = create_response.data["id"]

            response = self.client.post(
                f"/api/documents/{document_id}/reindex/",
            )

        self.assertEqual(
            response.status_code,
            status.HTTP_200_OK,
        )

        self.assertGreater(
            response.data["chunk_count"],
            0,
        )

    def test_delete_endpoint(self):
        with self.mocked_embedding_service():
            content = self.make_docx(
                paragraphs=[
                    "متن delete endpoint " * 100,
                ],
            )

            create_response = self.client.post(
                "/api/documents/",
                {
                    "title": "Delete API",
                    "file": self.make_upload(content),
                },
                format="multipart",
            )

        document_id = create_response.data["id"]

        response = self.client.delete(
            f"/api/documents/{document_id}/",
        )

        self.assertEqual(
            response.status_code,
            status.HTTP_204_NO_CONTENT,
        )

        self.assertFalse(
            Document.objects.filter(
                pk=document_id,
            ).exists()
        )

    def test_document_not_found(self):
        response = self.client.get(
            "/api/documents/999999/",
        )

        self.assertEqual(
            response.status_code,
            status.HTTP_404_NOT_FOUND,
        )

        self.assertEqual(
            response.data["error"]["code"],
            "not_found",
        )


class RealEmbeddingIntegrationTests(
    DocumentTestMixin,
    TemporaryMediaMixin,
    TestCase,
):
    """
    Verify the complete ingestion pipeline with the real
    SentenceTransformer model and PostgreSQL/pgvector.
    """

    @classmethod
    def setUpClass(cls):
        super().setUpClass()

        get_embedding_service.cache_clear()

    @classmethod
    def tearDownClass(cls):
        get_embedding_service.cache_clear()

        super().tearDownClass()

    def test_real_docx_to_embedding_to_pgvector(self):
        content = self.make_docx(
            paragraphs=[
                (
                    "این یک سند واقعی برای تست "
                    "سیستم Document Q&A است. "
                )
                * 150,
                (
                    "این بخش برای بررسی embedding "
                    "متن فارسی استفاده می‌شود. "
                )
                * 150,
            ],
            table_rows=[
                ["موضوع", "مقدار"],
                ["Embedding", "384"],
                ["Vector DB", "PostgreSQL + pgvector"],
            ],
        )

        uploaded_file = self.make_upload(
            content,
            name="real-integration-test.docx",
        )

        document = create_document(
            title="Real Embedding Integration Test",
            uploaded_file=uploaded_file,
        )

        self.assertTrue(
            document.full_text,
        )

        self.assertIn(
            "سیستم Document Q&A",
            document.full_text,
        )

        self.assertGreater(
            document.chunk_count,
            0,
        )

        self.assertEqual(
            document.chunk_count,
            document.chunks.count(),
        )

        chunks = list(
            document.chunks.order_by(
                "chunk_index",
            )
        )

        self.assertEqual(
            len(chunks),
            document.chunk_count,
        )

        for chunk in chunks:
            self.assertEqual(
                len(chunk.embedding),
                settings.EMBEDDING_DIMENSION,
            )

        first_vector = chunks[0].embedding

        norm = math.sqrt(
            sum(
                value * value
                for value in first_vector
            )
        )

        self.assertAlmostEqual(
            norm,
            1.0,
            places=3,
        )

        with connection.cursor() as cursor:
            cursor.execute(
                """
                SELECT
                    vector_dims(embedding),
                    COUNT(*)
                FROM documents_documentchunk
                WHERE document_id = %s
                GROUP BY vector_dims(embedding)
                """,
                [document.id],
            )

            result = cursor.fetchone()

        self.assertIsNotNone(
            result,
        )

        vector_dimension, chunk_count = result

        self.assertEqual(
            vector_dimension,
            settings.EMBEDDING_DIMENSION,
        )

        self.assertEqual(
            chunk_count,
            document.chunk_count,
        )

        self.assertEqual(
            document.embedding_model,
            settings.EMBEDDING_MODEL,
        )

        self.assertTrue(
            default_storage.exists(
                document.file.name,
            )
        )



class TokenizerTests(TestCase):
    def test_english_tokenization(self):
        tokens = tokenize_text("Hello, World! 123")
        self.assertEqual(
            tokens,
            ["hello", "world", "123"],
        )

    def test_persian_tokenization(self):
        tokens = tokenize_text("مرخصی کارکنان سالانه")
        self.assertIn("مرخصی", tokens)
        self.assertIn("کارکنان", tokens)
        self.assertIn("سالانه", tokens)

    def test_mixed_persian_english(self):
        tokens = tokenize_text("قرارداد Contract سالانه 2024")
        self.assertIn("قرارداد", tokens)
        self.assertIn("contract", tokens)
        self.assertIn("سالانه", tokens)
        self.assertIn("2024", tokens)

    def test_punctuation_ignored(self):
        tokens = tokenize_text("hello!!! world???")
        self.assertEqual(tokens, ["hello", "world"])

    def test_persian_character_normalization(self):
        # Arabic yeh/kaf should normalize to Persian forms
        arabic = tokenize_text("يك قرارداد")
        persian = tokenize_text("یک قرارداد")
        self.assertEqual(arabic, persian)


class ScoringTests(TestCase):
    def test_keyword_score_exact_overlap(self):
        score = compute_keyword_score(
            "مرخصی کارکنان",
            "مرخصی کارکنان سالانه بیست روز است",
        )
        self.assertEqual(score, 1.0)

    def test_keyword_score_no_overlap(self):
        score = compute_keyword_score(
            "مرخصی کارکنان",
            "ساعات کاری شرکت هشت صبح است",
        )
        self.assertEqual(score, 0.0)

    def test_keyword_score_partial_overlap(self):
        score = compute_keyword_score(
            "مرخصی سالانه کارکنان",
            "مرخصی کارکنان در شرکت",
        )
        self.assertAlmostEqual(score, 2 / 3, places=5)

    def test_vector_score_formula(self):
        self.assertAlmostEqual(
            compute_vector_score(0.0),
            1.0,
            places=6,
        )
        self.assertAlmostEqual(
            compute_vector_score(1.0),
            0.5,
            places=6,
        )


class RetrievalServiceTests(
    DocumentTestMixin,
    TemporaryMediaMixin,
    TestCase,
):
    def _seed_chunks(self):
        """
        Create three documents with one deterministic chunk each.
        Embeddings are mocked so values only need correct dimension.
        """
        docs = []
        texts = [
            ("مرخصی", "مرخصی کارکنان سالانه بیست روز است."),
            ("ساعات", "ساعات کاری شرکت از ساعت هشت صبح تا چهار بعدازظهر است."),
            ("قرارداد", "قرارداد کارکنان به صورت یک ساله تنظیم می‌شود."),
        ]

        for index, (title, chunk_text) in enumerate(texts):
            document = Document.objects.create(
                title=title,
                full_text=chunk_text,
                chunk_count=1,
                embedding_model=settings.EMBEDDING_MODEL,
            )
            # minimal valid file field for model constraints
            from django.core.files.base import ContentFile
            document.file.save(
                f"{title}.docx",
                ContentFile(b"dummy"),
                save=True,
            )

            # Different embeddings so distance ordering is stable under mock
            embedding = [0.0] * settings.EMBEDDING_DIMENSION
            embedding[0] = 1.0 if index == 0 else 0.0
            embedding[1] = 1.0 if index == 1 else 0.0
            embedding[2] = 1.0 if index == 2 else 0.0

            DocumentChunk.objects.create(
                document=document,
                chunk_index=0,
                text=chunk_text,
                embedding=embedding,
            )
            docs.append(document)

        return docs

    def test_empty_question_rejected(self):
        with self.assertRaises(DocumentRetrievalError) as ctx:
            retrieve_chunks("")
        self.assertEqual(ctx.exception.code, "empty_question")

    def test_whitespace_question_rejected(self):
        with self.assertRaises(DocumentRetrievalError) as ctx:
            retrieve_chunks("   ")
        self.assertEqual(ctx.exception.code, "empty_question")

    def test_top_k_zero_rejected(self):
        with self.assertRaises(DocumentRetrievalError) as ctx:
            retrieve_chunks("سوال", top_k=0)
        self.assertEqual(ctx.exception.code, "invalid_top_k")

    def test_top_k_negative_rejected(self):
        with self.assertRaises(DocumentRetrievalError) as ctx:
            retrieve_chunks("سوال", top_k=-1)
        self.assertEqual(ctx.exception.code, "invalid_top_k")

    def test_top_k_too_large_rejected(self):
        with self.assertRaises(DocumentRetrievalError) as ctx:
            retrieve_chunks("سوال", top_k=100)
        self.assertEqual(ctx.exception.code, "invalid_top_k")

    def test_empty_document_ids_rejected(self):
        with self.assertRaises(DocumentRetrievalError) as ctx:
            retrieve_chunks("سوال", document_ids=[])
        self.assertEqual(ctx.exception.code, "empty_document_ids")

    def test_retrieve_returns_empty_when_no_chunks(self):
        with self.mocked_embedding_service():
            results = retrieve_chunks("هر سوالی")
        self.assertEqual(results, [])

    def test_hybrid_ranking_and_metadata(self):
        docs = self._seed_chunks()

        # Question embedding aligned with "مرخصی" chunk (dim 0)
        with self.mocked_embedding_service(value=0.0) as embedding_service:
            embedding_service.embed_documents.side_effect = (
                lambda texts: [
                    [1.0, 0.0, 0.0] + [0.0] * (settings.EMBEDDING_DIMENSION - 3)
                    for _ in texts
                ]
            )

            results = retrieve_chunks(
                "میزان مرخصی سالانه کارکنان چقدر است؟",
                top_k=2,
            )

        self.assertEqual(len(results), 2)
        self.assertEqual(results[0].document_id, docs[0].id)
        self.assertEqual(results[0].document_title, "مرخصی")
        self.assertEqual(results[0].chunk_index, 0)
        self.assertIn("مرخصی", results[0].text)
        self.assertGreater(
            results[0].final_score,
            results[1].final_score,
        )

        # final_score formula
        expected = (
            0.7 * results[0].vector_score
            + 0.3 * results[0].keyword_score
        )
        self.assertAlmostEqual(
            results[0].final_score,
            expected,
            places=6,
        )

    def test_document_ids_filter(self):
        docs = self._seed_chunks()

        with self.mocked_embedding_service(value=0.01):
            results = retrieve_chunks(
                "ساعات کاری",
                top_k=6,
                document_ids=[docs[1].id],
            )

        self.assertTrue(results)
        for item in results:
            self.assertEqual(item.document_id, docs[1].id)

    def test_top_k_limits_results(self):
        docs = self._seed_chunks()

        with self.mocked_embedding_service(value=0.01):
            results = retrieve_chunks(
                "کارکنان",
                top_k=1,
            )

        self.assertEqual(len(results), 1)

    def test_uses_embedding_service(self):
        self._seed_chunks()

        with self.mocked_embedding_service() as embedding_service:
            retrieve_chunks("سوال آزمایشی", top_k=1)

        embedding_service.embed_documents.assert_called()


class RetrievalAPITests(
    DocumentTestMixin,
    TemporaryMediaMixin,
    APITestCase,
):
    def setUp(self):
        self.client = APIClient()

    def test_retrieval_endpoint_validation(self):
        response = self.client.post(
            "/api/retrieval/",
            {"question": "   "},
            format="json",
        )
        self.assertEqual(
            response.status_code,
            status.HTTP_400_BAD_REQUEST,
        )

    def test_retrieval_endpoint_success(self):
        document = Document.objects.create(
            title="Policy",
            full_text="مرخصی کارکنان سالانه بیست روز است.",
            chunk_count=1,
            embedding_model=settings.EMBEDDING_MODEL,
        )
        from django.core.files.base import ContentFile
        document.file.save(
            "policy.docx",
            ContentFile(b"dummy"),
            save=True,
        )
        embedding = [0.01] * settings.EMBEDDING_DIMENSION
        DocumentChunk.objects.create(
            document=document,
            chunk_index=0,
            text="مرخصی کارکنان سالانه بیست روز است.",
            embedding=embedding,
        )

        with self.mocked_embedding_service():
            response = self.client.post(
                "/api/retrieval/",
                {
                    "question": "مرخصی کارکنان چقدر است؟",
                    "top_k": 3,
                },
                format="json",
            )

        self.assertEqual(
            response.status_code,
            status.HTTP_200_OK,
        )
        self.assertEqual(
            response.data["question"],
            "مرخصی کارکنان چقدر است؟",
        )
        self.assertIsInstance(response.data["results"], list)
        if response.data["results"]:
            first = response.data["results"][0]
            self.assertIn("document_id", first)
            self.assertIn("final_score", first)
            self.assertIn("text", first)



class RAGContextTests(TestCase):
    def test_build_context_includes_metadata(self):
        chunks = [
            RetrievedChunk(
                document_id=1,
                document_title="Policy",
                chunk_index=0,
                text="مرخصی بیست روز است.",
                distance=0.1,
                vector_score=0.9,
                keyword_score=0.5,
                final_score=0.78,
            )
        ]
        context = build_rag_context(chunks)
        self.assertIn("document_id: 1", context)
        self.assertIn("document_title: Policy", context)
        self.assertIn("chunk_index: 0", context)
        self.assertIn("مرخصی بیست روز است.", context)
        self.assertNotIn("embedding", context.lower())

    def test_parse_valid_json(self):
        chunks = [
            RetrievedChunk(
                document_id=4,
                document_title="Company Policy",
                chunk_index=3,
                text="مدت قرارداد یک سال است.",
                distance=0.2,
                vector_score=0.8,
                keyword_score=0.4,
                final_score=0.68,
            )
        ]
        raw = (
            '{"answer": "یک سال", "sources": [{"document_id": 4, '
            '"document_title": "Company Policy", "chunk_index": 3, '
            '"excerpt": "مدت قرارداد یک سال است."}], "confidence": "high"}'
        )
        answer, sources, confidence = parse_llm_response(raw, chunks)
        self.assertEqual(answer, "یک سال")
        self.assertEqual(confidence, "high")
        self.assertEqual(len(sources), 1)
        self.assertEqual(sources[0].document_id, 4)
        self.assertLessEqual(len(sources[0].excerpt), 250)

    def test_parse_markdown_json(self):
        chunks = [
            RetrievedChunk(
                document_id=1,
                document_title="A",
                chunk_index=0,
                text="hello world",
                distance=0.1,
                vector_score=0.9,
                keyword_score=0.1,
                final_score=0.66,
            )
        ]
        raw = '```json\n{"answer": "hello", "sources": [], "confidence": "medium"}\n```'
        answer, sources, confidence = parse_llm_response(raw, chunks)
        self.assertEqual(answer, "hello")
        self.assertEqual(confidence, "medium")
        self.assertEqual(sources, [])

    def test_parse_invalid_json_fallback(self):
        chunks = []
        raw = "this is not json at all"
        answer, sources, confidence = parse_llm_response(raw, chunks)
        self.assertEqual(answer, raw)
        self.assertEqual(sources, [])
        self.assertEqual(confidence, "low")

    def test_invalid_confidence_becomes_low(self):
        chunks = []
        raw = '{"answer": "x", "sources": [], "confidence": "extreme"}'
        answer, sources, confidence = parse_llm_response(raw, chunks)
        self.assertEqual(confidence, "low")

    def test_invented_source_rejected(self):
        chunks = [
            RetrievedChunk(
                document_id=1,
                document_title="A",
                chunk_index=0,
                text="real text",
                distance=0.1,
                vector_score=0.9,
                keyword_score=0.2,
                final_score=0.69,
            )
        ]
        raw = (
            '{"answer": "x", "sources": [{"document_id": 999, '
            '"document_title": "Fake", "chunk_index": 0, '
            '"excerpt": "nope"}], "confidence": "high"}'
        )
        answer, sources, confidence = parse_llm_response(raw, chunks)
        self.assertEqual(sources, [])


class LLMConfigTests(TestCase):
    def tearDown(self):
        get_chat_model.cache_clear()
        super().tearDown()

    @override_settings(OPENROUTER_API_KEY="", OPENROUTER_MODEL="")
    def test_not_configured(self):
        self.assertFalse(is_llm_configured())

    @override_settings(
        OPENROUTER_API_KEY="test-key",
        OPENROUTER_MODEL="test-model",
    )
    def test_configured(self):
        self.assertTrue(is_llm_configured())

    @override_settings(OPENROUTER_API_KEY="", OPENROUTER_MODEL="m")
    def test_get_chat_model_missing_key(self):
        get_chat_model.cache_clear()
        with self.assertRaises(RAGError) as ctx:
            get_chat_model()
        self.assertEqual(ctx.exception.code, "llm_not_configured")


class QAServiceTests(
    DocumentTestMixin,
    TemporaryMediaMixin,
    TestCase,
):
    def tearDown(self):
        get_chat_model.cache_clear()
        super().tearDown()

    def _create_chunk_document(self, title, text, value=0.01):
        with self.mocked_embedding_service(value=value):
            content = self.make_docx(paragraphs=[text * 15])
            return create_document(
                title=title,
                uploaded_file=self.make_upload(content),
            )

    def test_no_context_does_not_call_llm(self):
        with self.mocked_embedding_service(value=0.02):
            with patch(
                "apps.documents.services.get_chat_model",
            ) as mock_get:
                result = answer_question(
                    "سؤالی که هیچ سندی ندارد",
                    top_k=3,
                )

        mock_get.assert_not_called()
        self.assertEqual(result.answer, NO_CONTEXT_ANSWER)
        self.assertEqual(result.sources, [])
        self.assertEqual(result.confidence, "low")

    def test_successful_qa(self):
        self._create_chunk_document(
            "Leave",
            "مرخصی کارکنان سالانه بیست روز است.",
        )

        mock_model = Mock()
        mock_response = Mock()
        mock_response.content = (
            '{"answer": "بیست روز", "sources": [{"document_id": 1, '
            '"document_title": "Leave", "chunk_index": 0, '
            '"excerpt": "مرخصی کارکنان سالانه بیست روز است."}], '
            '"confidence": "high"}'
        )
        mock_model.invoke.return_value = mock_response

        with self.mocked_embedding_service(value=0.02):
            with patch(
                "apps.documents.services.get_chat_model",
                return_value=mock_model,
            ):
                # Fix document_id in JSON dynamically after create
                doc = Document.objects.get(title="Leave")
                mock_response.content = (
                    '{"answer": "بیست روز", "sources": [{'
                    f'"document_id": {doc.id}, '
                    '"document_title": "Leave", "chunk_index": 0, '
                    '"excerpt": "مرخصی کارکنان سالانه بیست روز است."'
                    '}], "confidence": "high"}'
                )
                result = answer_question(
                    "میزان مرخصی سالانه چقدر است؟",
                    top_k=3,
                )

        self.assertEqual(result.answer, "بیست روز")
        self.assertEqual(result.confidence, "high")
        self.assertTrue(result.sources)
        self.assertEqual(result.sources[0].document_id, doc.id)
        mock_model.invoke.assert_called_once()

    def test_llm_failure_raises_rag_error(self):
        self._create_chunk_document(
            "X",
            "متن نمونه برای خطا.",
        )
        mock_model = Mock()
        mock_model.invoke.side_effect = RuntimeError("timeout")

        with self.mocked_embedding_service(value=0.02):
            with patch(
                "apps.documents.services.get_chat_model",
                return_value=mock_model,
            ):
                with self.assertRaises(RAGError) as ctx:
                    answer_question("متن نمونه چیست؟", top_k=2)

        self.assertEqual(ctx.exception.code, "llm_request_failed")


class QuestionAPITests(
    DocumentTestMixin,
    TemporaryMediaMixin,
    APITestCase,
):
    def setUp(self):
        self.client = APIClient()

    def tearDown(self):
        get_chat_model.cache_clear()
        super().tearDown()

    def test_question_endpoint_no_context(self):
        with self.mocked_embedding_service(value=0.02):
            with patch(
                "apps.documents.services.get_chat_model",
            ) as mock_get:
                response = self.client.post(
                    "/api/questions/",
                    {"question": "چیزی که در دیتابیس نیست"},
                    format="json",
                )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(
            response.data["answer"],
            NO_CONTEXT_ANSWER,
        )
        self.assertEqual(response.data["sources"], [])
        self.assertEqual(response.data["confidence"], "low")
        mock_get.assert_not_called()

    def test_question_endpoint_success(self):
        with self.mocked_embedding_service(value=0.02):
            content = self.make_docx(
                paragraphs=[
                    "ساعات کاری از هشت تا چهار است. " * 20,
                ],
            )
            create_response = self.client.post(
                "/api/documents/",
                {
                    "title": "Hours",
                    "file": self.make_upload(content),
                },
                format="multipart",
            )
            doc_id = create_response.data["id"]

            mock_model = Mock()
            mock_response = Mock()
            mock_response.content = (
                '{"answer": "از هشت تا چهار", "sources": [{'
                f'"document_id": {doc_id}, '
                '"document_title": "Hours", "chunk_index": 0, '
                '"excerpt": "ساعات کاری از هشت تا چهار است."'
                '}], "confidence": "medium"}'
            )
            mock_model.invoke.return_value = mock_response

            with patch(
                "apps.documents.services.get_chat_model",
                return_value=mock_model,
            ):
                response = self.client.post(
                    "/api/questions/",
                    {
                        "question": "ساعات کاری چیست؟",
                        "top_k": 3,
                    },
                    format="json",
                )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data["answer"], "از هشت تا چهار")
        self.assertEqual(response.data["confidence"], "medium")
        self.assertIn("latency_ms", response.data)
        self.assertTrue(response.data["sources"])

    def test_question_rejects_empty(self):
        response = self.client.post(
            "/api/questions/",
            {"question": "  "},
            format="json",
        )
        self.assertEqual(
            response.status_code,
            status.HTTP_400_BAD_REQUEST,
        )



class QuestionHistoryServiceTests(
    DocumentTestMixin,
    TemporaryMediaMixin,
    TestCase,
):
    def tearDown(self):
        get_chat_model.cache_clear()
        super().tearDown()

    def test_successful_qa_persists_question_and_answer(self):
        with self.mocked_embedding_service(value=0.02):
            content = self.make_docx(
                paragraphs=["مرخصی کارکنان سالانه بیست روز است. " * 20],
            )
            doc = create_document(
                title="Leave",
                uploaded_file=self.make_upload(content),
            )

            mock_model = Mock()
            mock_response = Mock()
            mock_response.content = (
                '{"answer": "بیست روز", "sources": [{'
                f'"document_id": {doc.id}, '
                '"document_title": "Leave", "chunk_index": 0, '
                '"excerpt": "مرخصی کارکنان سالانه بیست روز است."'
                '}], "confidence": "high"}'
            )
            mock_model.invoke.return_value = mock_response

            with patch(
                "apps.documents.services.get_chat_model",
                return_value=mock_model,
            ):
                question_row, result = ask_and_persist_question(
                    "میزان مرخصی چقدر است؟",
                    top_k=3,
                    document_ids=[doc.id],
                )

        self.assertEqual(Question.objects.count(), 1)
        self.assertEqual(Answer.objects.count(), 1)
        self.assertEqual(question_row.text, "میزان مرخصی چقدر است؟")
        self.assertEqual(question_row.top_k, 3)
        self.assertEqual(question_row.document_ids, [doc.id])

        answer = question_row.answer
        self.assertEqual(answer.text, "بیست روز")
        self.assertEqual(answer.confidence, "high")
        self.assertFalse(answer.error)
        self.assertTrue(answer.sources_json)
        self.assertLessEqual(
            len(answer.sources_json[0].get("excerpt", "")),
            250,
        )
        self.assertEqual(result.answer, "بیست روز")

    def test_llm_failure_keeps_question_and_error_answer(self):
        with self.mocked_embedding_service(value=0.02):
            content = self.make_docx(
                paragraphs=["متن نمونه برای خطا. " * 20],
            )
            create_document(
                title="X",
                uploaded_file=self.make_upload(content),
            )

            mock_model = Mock()
            mock_model.invoke.side_effect = RuntimeError("timeout")

            with patch(
                "apps.documents.services.get_chat_model",
                return_value=mock_model,
            ):
                with self.assertRaises(RAGError):
                    ask_and_persist_question(
                        "متن نمونه چیست؟",
                        top_k=2,
                    )

        self.assertEqual(Question.objects.count(), 1)
        self.assertEqual(Answer.objects.count(), 1)
        answer = Answer.objects.get()
        self.assertTrue(answer.error)
        self.assertEqual(answer.confidence, "low")
        self.assertEqual(answer.text, "")

    def test_validation_does_not_create_history(self):
        with self.assertRaises(DocumentRetrievalError):
            ask_and_persist_question("   ")

        self.assertEqual(Question.objects.count(), 0)
        self.assertEqual(Answer.objects.count(), 0)


class QuestionHistoryAPITests(
    DocumentTestMixin,
    TemporaryMediaMixin,
    APITestCase,
):
    def setUp(self):
        self.client = APIClient()

    def tearDown(self):
        get_chat_model.cache_clear()
        super().tearDown()

    def test_post_creates_history(self):
        with self.mocked_embedding_service(value=0.02):
            content = self.make_docx(
                paragraphs=["ساعات کاری از هشت تا چهار است. " * 20],
            )
            create_response = self.client.post(
                "/api/documents/",
                {
                    "title": "Hours",
                    "file": self.make_upload(content),
                },
                format="multipart",
            )
            doc_id = create_response.data["id"]

            mock_model = Mock()
            mock_response = Mock()
            mock_response.content = (
                '{"answer": "از هشت تا چهار", "sources": [{'
                f'"document_id": {doc_id}, '
                '"document_title": "Hours", "chunk_index": 0, '
                '"excerpt": "ساعات کاری از هشت تا چهار است."'
                '}], "confidence": "medium"}'
            )
            mock_model.invoke.return_value = mock_response

            with patch(
                "apps.documents.services.get_chat_model",
                return_value=mock_model,
            ):
                response = self.client.post(
                    "/api/questions/",
                    {
                        "question": "ساعات کاری چیست؟",
                        "top_k": 3,
                        "document_ids": [doc_id],
                    },
                    format="json",
                )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertIn("id", response.data)
        self.assertEqual(Question.objects.count(), 1)
        self.assertEqual(Answer.objects.count(), 1)
        q = Question.objects.get()
        self.assertEqual(q.document_ids, [doc_id])
        self.assertEqual(q.top_k, 3)

    def test_list_history(self):
        Question.objects.create(text="q1", top_k=6)
        Question.objects.create(text="q2", top_k=6)

        response = self.client.get("/api/questions/")
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertIn("results", response.data)
        self.assertEqual(len(response.data["results"]), 2)
        # newest first
        self.assertEqual(
            response.data["results"][0]["question"],
            "q2",
        )

    def test_detail_history(self):
        q = Question.objects.create(
            text="detail question",
            top_k=4,
            document_ids=[1, 2],
        )
        Answer.objects.create(
            question=q,
            text="detail answer",
            confidence="high",
            sources_json=[
                {
                    "document_id": 1,
                    "document_title": "D",
                    "chunk_index": 0,
                    "excerpt": "ex",
                }
            ],
            model_name="test-model",
            latency_ms=100,
        )

        response = self.client.get(f"/api/questions/{q.id}/")
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data["question"], "detail question")
        self.assertEqual(response.data["top_k"], 4)
        self.assertEqual(response.data["document_ids"], [1, 2])
        self.assertEqual(response.data["answer"]["text"], "detail answer")
        self.assertEqual(response.data["answer"]["confidence"], "high")

    def test_detail_not_found(self):
        response = self.client.get("/api/questions/999999/")
        self.assertEqual(response.status_code, status.HTTP_404_NOT_FOUND)
        self.assertEqual(
            response.data["error"]["code"],
            "question_not_found",
        )

    def test_post_validation_no_history(self):
        response = self.client.post(
            "/api/questions/",
            {"question": "  "},
            format="json",
        )
        self.assertEqual(
            response.status_code,
            status.HTTP_400_BAD_REQUEST,
        )
        self.assertEqual(Question.objects.count(), 0)

    def test_llm_failure_http_503_with_history(self):
        with self.mocked_embedding_service(value=0.02):
            content = self.make_docx(
                paragraphs=["متن. " * 30],
            )
            self.client.post(
                "/api/documents/",
                {
                    "title": "Doc",
                    "file": self.make_upload(content),
                },
                format="multipart",
            )

            mock_model = Mock()
            mock_model.invoke.side_effect = RuntimeError("down")

            with patch(
                "apps.documents.services.get_chat_model",
                return_value=mock_model,
            ):
                response = self.client.post(
                    "/api/questions/",
                    {"question": "متن چیست؟"},
                    format="json",
                )

        self.assertEqual(
            response.status_code,
            status.HTTP_503_SERVICE_UNAVAILABLE,
        )
        self.assertEqual(Question.objects.count(), 1)
        self.assertEqual(Answer.objects.count(), 1)
        self.assertTrue(Answer.objects.get().error)


class QuestionAdminRegistrationTests(TestCase):
    def test_question_and_answer_registered(self):
        from django.contrib import admin as dj_admin
        from .models import Answer, Question

        self.assertTrue(dj_admin.site.is_registered(Question))
        self.assertTrue(dj_admin.site.is_registered(Answer))


class APIDocumentationTests(APITestCase):
    """Milestone 6: schema, routing, pagination, error envelope."""

    def setUp(self):
        self.client = APIClient()

    def test_health_endpoint(self):
        response = self.client.get("/api/health/")
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        for key in ("status", "db", "vector", "llm_configured"):
            self.assertIn(key, response.data)

    def test_openapi_schema_available(self):
        response = self.client.get("/api/schema/")
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        content = response.content.decode("utf-8", errors="ignore")
        self.assertTrue(
            "openapi" in content.lower()
            or "Document Q&A" in content
            or "paths" in content
        )

    def test_swagger_ui_available(self):
        response = self.client.get("/api/docs/")
        self.assertEqual(response.status_code, status.HTTP_200_OK)

    def test_documents_list_pagination_shape(self):
        response = self.client.get("/api/documents/")
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertIn("count", response.data)
        self.assertIn("results", response.data)
        self.assertIn("next", response.data)
        self.assertIn("previous", response.data)

    def test_questions_list_pagination_shape(self):
        response = self.client.get("/api/questions/")
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertIn("count", response.data)
        self.assertIn("results", response.data)

    def test_document_not_found_error_envelope(self):
        response = self.client.get("/api/documents/999999/")
        self.assertEqual(response.status_code, status.HTTP_404_NOT_FOUND)
        self.assertIn("error", response.data)
        self.assertIn("code", response.data["error"])
        self.assertIn("message", response.data["error"])

    def test_question_not_found_error_envelope(self):
        response = self.client.get("/api/questions/999999/")
        self.assertEqual(response.status_code, status.HTTP_404_NOT_FOUND)
        self.assertEqual(
            response.data["error"]["code"],
            "question_not_found",
        )

    def test_schema_mentions_core_paths(self):
        response = self.client.get(
            "/api/schema/",
            HTTP_ACCEPT="application/json",
        )
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        data = response.data if isinstance(response.data, dict) else None
        if data is None:
            import json
            data = json.loads(response.content.decode())
        paths = data.get("paths", {})
        path_keys = " ".join(paths.keys())
        self.assertIn("/api/documents/", path_keys)
        self.assertIn("/api/questions/", path_keys)
        self.assertIn("/api/retrieval/", path_keys)
        self.assertIn("/api/health/", path_keys)

