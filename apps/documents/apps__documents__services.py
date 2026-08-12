import logging
import re
import unicodedata
import zipfile
from dataclasses import dataclass
from functools import lru_cache
from pathlib import Path
from typing import BinaryIO

from django.conf import settings
from django.core.files.storage import default_storage
from django.db import transaction
from docx import Document as DocxDocument
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph
from langchain_text_splitters import RecursiveCharacterTextSplitter
from pgvector.django import CosineDistance
from sentence_transformers import SentenceTransformer

from .models import Document, DocumentChunk

logger = logging.getLogger(__name__)


class DocumentIngestionError(Exception):
    """Raised when document ingestion cannot be completed."""

    def __init__(
        self,
        message: str,
        code: str = "document_ingestion_error",
    ) -> None:
        super().__init__(message)
        self.message = message
        self.code = code


class EmbeddingConfigurationError(DocumentIngestionError):
    """Raised when the embedding model is not configured correctly."""


@dataclass(frozen=True)
class PreparedDocument:
    full_text: str
    chunks: list[str]
    embeddings: list[list[float]]
    embedding_model: str


def _max_upload_size_bytes() -> int:
    return settings.MAX_UPLOAD_SIZE_MB * 1024 * 1024


def validate_docx_upload(uploaded_file: BinaryIO) -> None:
    """
    Validate the uploaded file before persistence.
    """
    filename = Path(
        getattr(uploaded_file, "name", "")
    ).name

    if not filename:
        raise DocumentIngestionError(
            "نام فایل ارسال نشده است.",
            "missing_filename",
        )

    if Path(filename).suffix.lower() != ".docx":
        raise DocumentIngestionError(
            "فقط فایل‌های DOCX مجاز هستند.",
            "invalid_extension",
        )

    file_size = getattr(uploaded_file, "size", None)

    if file_size is None:
        current_position = uploaded_file.tell()
        uploaded_file.seek(0, 2)
        file_size = uploaded_file.tell()
        uploaded_file.seek(current_position)

    if file_size == 0:
        raise DocumentIngestionError(
            "فایل خالی است.",
            "empty_file",
        )

    if file_size > _max_upload_size_bytes():
        raise DocumentIngestionError(
            (
                "حجم فایل بیشتر از "
                f"{settings.MAX_UPLOAD_SIZE_MB} مگابایت است."
            ),
            "file_too_large",
        )

    try:
        uploaded_file.seek(0)

        with zipfile.ZipFile(uploaded_file) as archive:
            names = set(archive.namelist())

            required_files = {
                "[Content_Types].xml",
                "word/document.xml",
            }

            if not required_files.issubset(names):
                raise DocumentIngestionError(
                    "ساختار فایل DOCX معتبر نیست.",
                    "invalid_docx_structure",
                )

            if any(
                info.flag_bits & 0x1
                for info in archive.infolist()
            ):
                raise DocumentIngestionError(
                    "فایل DOCX رمزگذاری‌شده قابل پردازش نیست.",
                    "encrypted_docx",
                )

            uncompressed_size = sum(
                info.file_size
                for info in archive.infolist()
            )

            max_uncompressed_size = (
                _max_upload_size_bytes() * 10
            )

            if uncompressed_size > max_uncompressed_size:
                raise DocumentIngestionError(
                    "ساختار فشرده فایل بیش از حد بزرگ است.",
                    "document_expansion_too_large",
                )

            corrupted_file = archive.testzip()

            if corrupted_file is not None:
                raise DocumentIngestionError(
                    "فایل DOCX خراب است.",
                    "corrupted_docx",
                )

    except DocumentIngestionError:
        raise
    except (zipfile.BadZipFile, OSError) as exc:
        raise DocumentIngestionError(
            "فایل DOCX معتبر نیست.",
            "invalid_docx",
        ) from exc
    finally:
        uploaded_file.seek(0)


def _iter_block_items(document: DocxDocument):
    """
    Yield paragraphs and tables in document order.
    """
    body = document.element.body

    for child in body.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, document)
        elif isinstance(child, CT_Tbl):
            yield Table(child, document)


def _normalize_text(text: str) -> str:
    """
    Normalize whitespace while preserving Persian characters and ZWNJ.
    """
    text = unicodedata.normalize("NFC", text)

    text = (
        text.replace("\r\n", "\n")
        .replace("\r", "\n")
    )

    text = re.sub(
        r"[\x00-\x08\x0B\x0C\x0E-\x1F\x7F]",
        "",
        text,
    )

    text = re.sub(
        r"[ \t]+",
        " ",
        text,
    )

    return text.strip()


def extract_docx_text(uploaded_file: BinaryIO) -> str:
    """
    Extract paragraphs and table rows from a DOCX file.
    """
    validate_docx_upload(uploaded_file)

    try:
        uploaded_file.seek(0)

        document = DocxDocument(uploaded_file)

        blocks: list[str] = []

        for block in _iter_block_items(document):
            if isinstance(block, Paragraph):
                text = _normalize_text(block.text)

                if text:
                    blocks.append(text)

            elif isinstance(block, Table):
                for row in block.rows:
                    cells = [
                        _normalize_text(cell.text)
                        for cell in row.cells
                    ]

                    cells = [
                        cell
                        for cell in cells
                        if cell
                    ]

                    if cells:
                        blocks.append(
                            " | ".join(cells)
                        )

        full_text = "\n\n".join(blocks).strip()

        if not full_text:
            raise DocumentIngestionError(
                "سند هیچ متن قابل پردازشی ندارد.",
                "empty_document",
            )

        return full_text

    except DocumentIngestionError:
        raise
    except Exception as exc:
        logger.exception("DOCX extraction failed")
        raise DocumentIngestionError(
            "استخراج متن فایل DOCX با خطا مواجه شد.",
            "docx_extraction_failed",
        ) from exc
    finally:
        uploaded_file.seek(0)


def split_document_text(full_text: str) -> list[str]:
    """
    Split document text using LangChain's recursive character splitter.
    """
    chunk_size = settings.CHUNK_SIZE
    chunk_overlap = settings.CHUNK_OVERLAP

    if chunk_size <= 0:
        raise DocumentIngestionError(
            "CHUNK_SIZE باید بزرگ‌تر از صفر باشد.",
            "invalid_chunk_size",
        )

    if chunk_overlap < 0 or chunk_overlap >= chunk_size:
        raise DocumentIngestionError(
            "CHUNK_OVERLAP باید بین صفر و CHUNK_SIZE باشد.",
            "invalid_chunk_overlap",
        )

    splitter = RecursiveCharacterTextSplitter(
        chunk_size=chunk_size,
        chunk_overlap=chunk_overlap,
        separators=[
            "\n\n",
            "\n",
            " ",
            "",
        ],
        length_function=len,
    )

    chunks = [
        chunk.strip()
        for chunk in splitter.split_text(full_text)
        if chunk.strip()
    ]

    if not chunks:
        raise DocumentIngestionError(
            "هیچ chunk قابل ذخیره‌ای از سند تولید نشد.",
            "empty_chunks",
        )

    return chunks


class EmbeddingService:
    """
    Provide local Sentence Transformer embeddings.
    """

    def __init__(self) -> None:
        model_name = settings.EMBEDDING_MODEL

        try:
            self.model = SentenceTransformer(
                model_name,
            )
        except Exception as exc:
            logger.exception(
                "Failed to load embedding model: %s",
                model_name,
            )

            raise EmbeddingConfigurationError(
                "مدل embedding قابل بارگذاری نیست.",
                "embedding_model_load_failed",
            ) from exc

        actual_dimension = (
            self.model.get_embedding_dimension()
        )

        expected_dimension = (
            settings.EMBEDDING_DIMENSION
        )

        if actual_dimension != expected_dimension:
            raise EmbeddingConfigurationError(
                (
                    "ابعاد مدل embedding با تنظیمات پروژه "
                    "مطابقت ندارد."
                ),
                "embedding_dimension_mismatch",
            )

        self.model_name = model_name

    def embed_documents(
        self,
        texts: list[str],
    ) -> list[list[float]]:
        """
        Encode multiple texts into normalized vectors.
        """
        if not texts:
            return []

        try:
            embeddings = self.model.encode(
                inputs=texts,
                batch_size=32,
                show_progress_bar=False,
                convert_to_numpy=True,
                normalize_embeddings=True,
            )
        except Exception as exc:
            logger.exception("Embedding generation failed")

            raise DocumentIngestionError(
                "تولید embedding با خطا مواجه شد.",
                "embedding_generation_failed",
            ) from exc

        if embeddings.ndim != 2:
            raise EmbeddingConfigurationError(
                "خروجی embedding معتبر نیست.",
                "invalid_embedding_shape",
            )

        if embeddings.shape[1] != (
            settings.EMBEDDING_DIMENSION
        ):
            raise EmbeddingConfigurationError(
                "ابعاد embedding با schema پایگاه داده یکسان نیست.",
                "embedding_dimension_mismatch",
            )

        return embeddings.astype(
            "float32"
        ).tolist()


@lru_cache(maxsize=1)
def get_embedding_service() -> EmbeddingService:
    """
    Return a singleton embedding service.
    """
    return EmbeddingService()


def prepare_document(
    uploaded_file: BinaryIO,
) -> PreparedDocument:
    """
    Validate, extract, chunk, and embed a document.
    """
    full_text = extract_docx_text(
        uploaded_file,
    )

    chunks = split_document_text(
        full_text,
    )

    embeddings = get_embedding_service().embed_documents(
        chunks,
    )

    return PreparedDocument(
        full_text=full_text,
        chunks=chunks,
        embeddings=embeddings,
        embedding_model=settings.EMBEDDING_MODEL,
    )


def _replace_chunks(
    document: Document,
    prepared: PreparedDocument,
) -> None:
    """
    Replace all existing chunks for a document.
    """
    DocumentChunk.objects.filter(
        document=document,
    ).delete()

    DocumentChunk.objects.bulk_create(
        [
            DocumentChunk(
                document=document,
                chunk_index=index,
                text=text,
                embedding=embedding,
            )
            for index, (text, embedding) in enumerate(
                zip(
                    prepared.chunks,
                    prepared.embeddings,
                    strict=True,
                )
            )
        ]
    )


def _delete_stored_file(file_name: str, full_path: str | None = None) -> None:
    """
    Delete a stored upload by storage name and/or absolute filesystem path.
    """
    if file_name:
        try:
            default_storage.delete(file_name)
        except Exception:
            logger.exception(
                "default_storage.delete failed",
                extra={"file_name": file_name},
            )

    candidates: list[Path] = []

    if full_path:
        candidates.append(Path(full_path))

    if file_name:
        candidates.append(Path(str(settings.MEDIA_ROOT)) / file_name)
        try:
            candidates.append(Path(default_storage.path(file_name)))
        except Exception:
            pass

    for path in candidates:
        try:
            if path.is_file():
                path.unlink()
        except Exception:
            logger.exception(
                "Filesystem unlink failed",
                extra={"path": str(path)},
            )


def create_document(
    *,
    title: str,
    uploaded_file,
) -> Document:
    """
    Validate, persist, extract, chunk, embed, and store a new document.
    """
    prepared = prepare_document(uploaded_file)

    saved_file_name: str | None = None
    saved_full_path: str | None = None

    try:
        with transaction.atomic():
            document = Document(
                title=title,
                full_text=prepared.full_text,
                chunk_count=len(prepared.chunks),
                embedding_model=prepared.embedding_model,
            )

            document.file.save(
                uploaded_file.name,
                uploaded_file,
                save=False,
            )

            saved_file_name = document.file.name

            try:
                saved_full_path = document.file.path
            except Exception:
                saved_full_path = str(
                    Path(str(settings.MEDIA_ROOT)) / saved_file_name
                )

            document.save()
            _replace_chunks(document, prepared)

    except Exception:
        _delete_stored_file(saved_file_name or "", saved_full_path)
        raise

    return document


def update_document(
    document: Document,
    *,
    title: str,
    uploaded_file=None,
) -> Document:
    """
    Update title only, or replace the source file and re-index.
    """
    if uploaded_file is None:
        document.title = title
        document.save(
            update_fields=[
                "title",
                "updated_at",
            ],
        )
        return document

    prepared = prepare_document(uploaded_file)

    old_file_name = document.file.name if document.file else None
    old_full_path: str | None = None

    if document.file:
        try:
            old_full_path = document.file.path
        except Exception:
            if old_file_name:
                old_full_path = str(
                    Path(str(settings.MEDIA_ROOT)) / old_file_name
                )

    new_file_name: str | None = None
    new_full_path: str | None = None

    try:
        with transaction.atomic():
            document.title = title
            document.full_text = prepared.full_text
            document.chunk_count = len(prepared.chunks)
            document.embedding_model = prepared.embedding_model

            document.file.save(
                uploaded_file.name,
                uploaded_file,
                save=False,
            )

            new_file_name = document.file.name

            try:
                new_full_path = document.file.path
            except Exception:
                new_full_path = str(
                    Path(str(settings.MEDIA_ROOT)) / new_file_name
                )

            document.save()
            _replace_chunks(document, prepared)

    except Exception:
        _delete_stored_file(new_file_name or "", new_full_path)
        raise

    if (
        old_file_name
        and old_file_name != new_file_name
    ):
        _delete_stored_file(old_file_name, old_full_path)

    return document


def reindex_document(
    document: Document,
) -> Document:
    """
    Re-extract and re-index the document's existing source file.
    """
    with transaction.atomic():
        locked_document = (
            Document.objects
            .select_for_update()
            .get(pk=document.pk)
        )

        try:
            with locked_document.file.open("rb") as file_obj:
                prepared = prepare_document(
                    file_obj,
                )
        except FileNotFoundError as exc:
            raise DocumentIngestionError(
                "فایل اصلی سند پیدا نشد.",
                "source_file_missing",
            ) from exc

        locked_document.full_text = (
            prepared.full_text
        )
        locked_document.chunk_count = (
            len(prepared.chunks)
        )
        locked_document.embedding_model = (
            prepared.embedding_model
        )

        locked_document.save(
            update_fields=[
                "full_text",
                "chunk_count",
                "embedding_model",
                "updated_at",
            ]
        )

        _replace_chunks(
            locked_document,
            prepared,
        )

    return locked_document


def delete_document(document: Document) -> None:
    file_name = document.file.name if document.file else None
    full_path: str | None = None

    if document.file:
        try:
            full_path = document.file.path
        except Exception:
            if file_name:
                full_path = str(
                    Path(str(settings.MEDIA_ROOT)) / file_name
                )

    with transaction.atomic():
        document.delete()

    if file_name:
        _delete_stored_file(file_name, full_path)


# ---------------------------------------------------------------------------
# Milestone 3 — Retrieval
# ---------------------------------------------------------------------------

MAX_RETRIEVER_TOP_K = 50

VECTOR_SCORE_WEIGHT = 0.7
KEYWORD_SCORE_WEIGHT = 0.3


class RetrievalError(Exception):
    """Raised when retrieval cannot be completed."""

    def __init__(
        self,
        message: str,
        code: str = "retrieval_error",
    ) -> None:
        super().__init__(message)
        self.message = message
        self.code = code


@dataclass(frozen=True)
class RetrievedChunk:
    document_id: int
    document_title: str
    chunk_index: int
    text: str
    distance: float
    vector_score: float
    keyword_score: float
    final_score: float


_PERSIAN_CHAR_MAP = str.maketrans(
    {
        "ي": "ی",
        "ك": "ک",
        "ة": "ه",
        "أ": "ا",
        "إ": "ا",
        "آ": "ا",
    }
)


def tokenize_text(text: str) -> list[str]:
    """
    Simple Unicode-aware tokenizer for Persian and English.
    """
    if not text:
        return []

    text = unicodedata.normalize("NFC", text)
    text = text.translate(_PERSIAN_CHAR_MAP)
    text = text.casefold()
    text = re.sub(r"[^\w]+", " ", text, flags=re.UNICODE)
    return [token for token in text.split() if token]


def compute_keyword_score(
    question_terms: list[str],
    chunk_text: str,
) -> float:
    """
    keyword_score = overlap / max(1, len(question_terms))
    """
    if not question_terms:
        return 0.0

    chunk_terms = set(tokenize_text(chunk_text))
    overlap = sum(
        1
        for term in question_terms
        if term in chunk_terms
    )
    return overlap / max(1, len(question_terms))


def compute_vector_score(distance: float) -> float:
    """
    vector_score = 1 / (1 + distance)
    """
    return 1.0 / (1.0 + float(distance))


def compute_final_score(
    vector_score: float,
    keyword_score: float,
) -> float:
    return (
        VECTOR_SCORE_WEIGHT * vector_score
        + KEYWORD_SCORE_WEIGHT * keyword_score
    )


def _validate_retrieval_question(question: str) -> str:
    if not isinstance(question, str):
        raise RetrievalError(
            "سؤال باید متن باشد.",
            "invalid_question",
        )

    question = question.strip()

    if not question:
        raise RetrievalError(
            "سؤال نمی‌تواند خالی باشد.",
            "empty_question",
        )

    return question


def _validate_top_k(top_k: int | None) -> int:
    if top_k is None:
        return settings.RETRIEVER_TOP_K

    if not isinstance(top_k, int) or isinstance(top_k, bool):
        raise RetrievalError(
            "مقدار top_k باید عدد صحیح باشد.",
            "invalid_top_k",
        )

    if top_k <= 0:
        raise RetrievalError(
            "مقدار top_k باید بزرگ‌تر از صفر باشد.",
            "invalid_top_k",
        )

    if top_k > MAX_RETRIEVER_TOP_K:
        raise RetrievalError(
            (
                f"مقدار top_k نمی‌تواند بیشتر از "
                f"{MAX_RETRIEVER_TOP_K} باشد."
            ),
            "invalid_top_k",
        )

    return top_k


def _validate_document_ids(
    document_ids: list[int] | None,
) -> list[int] | None:
    if document_ids is None:
        return None

    if not isinstance(document_ids, list):
        raise RetrievalError(
            "فیلتر document_ids باید لیست باشد.",
            "invalid_document_ids",
        )

    if len(document_ids) == 0:
        raise RetrievalError(
            "فیلتر document_ids نمی‌تواند خالی باشد.",
            "empty_document_ids",
        )

    cleaned: list[int] = []

    for value in document_ids:
        if not isinstance(value, int) or isinstance(value, bool):
            raise RetrievalError(
                "شناسه‌های سند باید عدد صحیح باشند.",
                "invalid_document_ids",
            )

        if value <= 0:
            raise RetrievalError(
                "شناسه‌های سند باید مثبت باشند.",
                "invalid_document_ids",
            )

        cleaned.append(value)

    return cleaned


def retrieve_chunks(
    *,
    question: str,
    top_k: int | None = None,
    document_ids: list[int] | None = None,
) -> list[RetrievedChunk]:
    """
    Hybrid retrieval: pgvector cosine search + lexical keyword rerank.
    """
    question = _validate_retrieval_question(question)
    top_k = _validate_top_k(top_k)
    document_ids = _validate_document_ids(document_ids)

    candidate_multiplier = (
        settings.RETRIEVER_CANDIDATE_MULTIPLIER
    )

    if candidate_multiplier < 1:
        raise RetrievalError(
            "ضریب کاندید بازیابی نامعتبر است.",
            "invalid_candidate_multiplier",
        )

    candidate_count = top_k * candidate_multiplier

    embeddings = get_embedding_service().embed_documents(
        [question],
    )

    if not embeddings:
        raise RetrievalError(
            "تولید embedding برای سؤال ناموفق بود.",
            "question_embedding_failed",
        )

    question_embedding = embeddings[0]

    if len(question_embedding) != settings.EMBEDDING_DIMENSION:
        raise RetrievalError(
            "ابعاد embedding سؤال با تنظیمات پروژه مطابقت ندارد.",
            "embedding_dimension_mismatch",
        )

    queryset = DocumentChunk.objects.select_related(
        "document",
    )

    if document_ids is not None:
        queryset = queryset.filter(
            document_id__in=document_ids,
        )

    candidates = list(
        queryset.annotate(
            distance=CosineDistance(
                "embedding",
                question_embedding,
            ),
        )
        .order_by("distance")[:candidate_count]
    )

    if not candidates:
        return []

    question_terms = tokenize_text(question)
    ranked: list[RetrievedChunk] = []

    for chunk in candidates:
        distance = float(chunk.distance)
        vector_score = compute_vector_score(distance)
        keyword_score = compute_keyword_score(
            question_terms,
            chunk.text,
        )
        final_score = compute_final_score(
            vector_score,
            keyword_score,
        )

        ranked.append(
            RetrievedChunk(
                document_id=chunk.document_id,
                document_title=chunk.document.title,
                chunk_index=chunk.chunk_index,
                text=chunk.text,
                distance=distance,
                vector_score=vector_score,
                keyword_score=keyword_score,
                final_score=final_score,
            )
        )

    ranked.sort(
        key=lambda item: item.final_score,
        reverse=True,
    )

    return ranked[:top_k]
