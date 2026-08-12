import math
import os
import shutil
import tempfile
from contextlib import contextmanager
from io import BytesIO
from pathlib import Path
from unittest.mock import Mock, patch

from rest_framework.test import (
    APIClient,
    APITestCase,
)
import numpy as np

from django.conf import settings
from django.core.files.storage import default_storage
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import connection
from django.test import (
    TestCase,
    override_settings,
)
from docx import Document as DocxDocument
from rest_framework import status

from .models import Document, DocumentChunk
from .services import (
    DocumentIngestionError,
    EmbeddingConfigurationError,
    EmbeddingService,
    RetrievalError,
    compute_final_score,
    compute_keyword_score,
    compute_vector_score,
    create_document,
    delete_document,
    extract_docx_text,
    get_embedding_service,
    reindex_document,
    retrieve_chunks,
    split_document_text,
    tokenize_text,
    update_document,
    validate_docx_upload,
)


class DocumentTestMixin:
    def make_docx(
        self,
        paragraphs: list[str],
        table_rows: list[list[str]] | None = None,
    ) -> bytes:
        document = DocxDocument()

        for paragraph in paragraphs:
            document.add_paragraph(paragraph)

        if table_rows:
            table = document.add_table(
                rows=1,
                cols=len(table_rows[0]),
            )

            for row_index, row in enumerate(table_rows):
                if row_index == 0:
                    cells = table.rows[0].cells
                else:
                    cells = table.add_row().cells

                for cell, value in zip(
                    cells,
                    row,
                    strict=True,
                ):
                    cell.text = value

        buffer = BytesIO()
        document.save(buffer)

        return buffer.getvalue()

    def make_upload(
        self,
        content: bytes,
        name: str = "sample.docx",
    ) -> SimpleUploadedFile:
        return SimpleUploadedFile(
            name,
            content,
            content_type=(
                "application/"
                "vnd.openxmlformats-officedocument."
                "wordprocessingml.document"
            ),
        )

    @contextmanager
    def mocked_embedding_service(
        self,
        model_name: str = "test-model",
        value: float = 0.01,
    ):
        embedding_service = Mock()
        embedding_service.model_name = model_name

        embedding_service.embed_documents.side_effect = (
            lambda texts: [
                [value] * settings.EMBEDDING_DIMENSION
                for _ in texts
            ]
        )

        with patch(
            "apps.documents.services.get_embedding_service",
            return_value=embedding_service,
        ):
            yield embedding_service


class TemporaryMediaMixin:
    @classmethod
    def setUpClass(cls):
        cls._media_root = Path(
            tempfile.mkdtemp(
                prefix="doc_qa_test_media_",
            )
        )

        cls._media_override = override_settings(
            MEDIA_ROOT=str(cls._media_root),
        )
        cls._media_override.enable()

        # Django caches storage instances and LazyObject wrappers.
        # Reset both so the new MEDIA_ROOT is used for file I/O.
        from django.core.files.storage import (
            default_storage,
            storages,
        )
        from django.utils.functional import empty

        storages._storages.clear()
        default_storage._wrapped = empty

        super().setUpClass()

    @classmethod
    def tearDownClass(cls):
        try:
            super().tearDownClass()
        finally:
            cls._media_override.disable()
            from django.core.files.storage import (
                default_storage,
                storages,
            )
            from django.utils.functional import empty

            storages._storages.clear()
            default_storage._wrapped = empty
            shutil.rmtree(
                cls._media_root,
                ignore_errors=True,
            )


class DocumentValidationTests(
    DocumentTestMixin,
    TestCase,
):
    def test_rejects_non_docx_extension(self):
        uploaded_file = self.make_upload(
            b"not a docx",
            name="sample.pdf",
        )

        with self.assertRaises(DocumentIngestionError) as context:
            validate_docx_upload(uploaded_file)

        self.assertEqual(
            context.exception.code,
            "invalid_extension",
        )

    def test_rejects_empty_file(self):
        uploaded_file = self.make_upload(
            b"",
        )

        with self.assertRaises(DocumentIngestionError) as context:
            validate_docx_upload(uploaded_file)

        self.assertEqual(
            context.exception.code,
            "empty_file",
        )

    def test_rejects_oversized_file(self):
        content = b"x" * (
            settings.MAX_UPLOAD_SIZE_MB * 1024 * 1024 + 1
        )

        uploaded_file = self.make_upload(
            content,
        )

        with self.assertRaises(DocumentIngestionError) as context:
            validate_docx_upload(uploaded_file)

        self.assertEqual(
            context.exception.code,
            "file_too_large",
        )

    def test_rejects_invalid_docx_structure(self):
        buffer = BytesIO()

        import zipfile

        with zipfile.ZipFile(
            buffer,
            "w",
            compression=zipfile.ZIP_DEFLATED,
        ) as archive:
            archive.writestr(
                "[Content_Types].xml",
                "<Types></Types>",
            )

        uploaded_file = self.make_upload(
            buffer.getvalue(),
        )

        with self.assertRaises(DocumentIngestionError) as context:
            validate_docx_upload(uploaded_file)

        self.assertEqual(
            context.exception.code,
            "invalid_docx_structure",
        )

    def test_rejects_corrupted_docx(self):
        uploaded_file = self.make_upload(
            b"this is not a zip file",
        )

        with self.assertRaises(DocumentIngestionError) as context:
            validate_docx_upload(uploaded_file)

        self.assertEqual(
            context.exception.code,
            "invalid_docx",
        )

    def test_validation_preserves_file_position(self):
        content = self.make_docx(
            paragraphs=["متن آزمایشی"],
        )

        uploaded_file = self.make_upload(content)

        uploaded_file.seek(5)

        validate_docx_upload(uploaded_file)

        self.assertEqual(
            uploaded_file.tell(),
            0,
        )


class DocumentExtractionTests(
    DocumentTestMixin,
    TestCase,
):
    def test_extracts_paragraphs_and_tables(self):
        content = self.make_docx(
            paragraphs=[
                "اولین پاراگراف",
                "پاراگراف دوم",
            ],
            table_rows=[
                ["نام", "مقدار"],
                ["A", "100"],
            ],
        )

        uploaded_file = self.make_upload(content)

        text = extract_docx_text(
            uploaded_file,
        )

        self.assertIn(
            "اولین پاراگراف",
            text,
        )

        self.assertIn(
            "پاراگراف دوم",
            text,
        )

        self.assertIn(
            "نام | مقدار",
            text,
        )

        self.assertIn(
            "A | 100",
            text,
        )

    def test_rejects_empty_document(self):
        document = DocxDocument()

        buffer = BytesIO()
        document.save(buffer)

        uploaded_file = self.make_upload(
            buffer.getvalue(),
        )

        with self.assertRaises(DocumentIngestionError) as context:
            extract_docx_text(uploaded_file)

        self.assertEqual(
            context.exception.code,
            "empty_document",
        )

    def test_normalizes_document_text(self):
        content = self.make_docx(
            paragraphs=[
                "  متن    با فاصله‌های زیاد  ",
            ],
        )

        uploaded_file = self.make_upload(content)

        text = extract_docx_text(
            uploaded_file,
        )

        self.assertEqual(
            text,
            "متن با فاصله‌های زیاد",
        )


class ChunkingTests(TestCase):
    def test_chunking_removes_empty_chunks(self):
        text = (
            "این یک متن آزمایشی است. " * 300
        )

        chunks = split_document_text(text)

        self.assertTrue(chunks)

        self.assertTrue(
            all(
                chunk.strip()
                for chunk in chunks
            )
        )

    def test_chunking_respects_configured_size(self):
        text = "A" * 3500

        chunks = split_document_text(text)

        self.assertTrue(chunks)

        self.assertTrue(
            all(
                len(chunk) <= settings.CHUNK_SIZE
                for chunk in chunks
            )
        )

    def test_chunking_uses_overlap(self):
        text = "A" * 2500

        chunks = split_document_text(text)

        self.assertGreaterEqual(
            len(chunks),
            2,
        )

        self.assertEqual(
            chunks[0][-settings.CHUNK_OVERLAP:],
            chunks[1][:settings.CHUNK_OVERLAP],
        )

    @override_settings(
        CHUNK_SIZE=0,
    )
    def test_rejects_invalid_chunk_size(self):
        with self.assertRaises(DocumentIngestionError) as context:
            split_document_text("some text")

        self.assertEqual(
            context.exception.code,
            "invalid_chunk_size",
        )

    @override_settings(
        CHUNK_SIZE=1000,
        CHUNK_OVERLAP=1000,
    )
    def test_rejects_invalid_chunk_overlap(self):
        with self.assertRaises(DocumentIngestionError) as context:
            split_document_text("some text")

        self.assertEqual(
            context.exception.code,
            "invalid_chunk_overlap",
        )


class EmbeddingServiceTests(TestCase):
    def tearDown(self):
        get_embedding_service.cache_clear()
        super().tearDown()

    @patch("apps.documents.services.SentenceTransformer")
    def test_loads_configured_model(
        self,
        mock_sentence_transformer,
    ):
        model = Mock()

        model.get_embedding_dimension.return_value = (
            settings.EMBEDDING_DIMENSION
        )

        mock_sentence_transformer.return_value = model

        service = EmbeddingService()

        mock_sentence_transformer.assert_called_once_with(
            settings.EMBEDDING_MODEL,
        )

        self.assertEqual(
            service.model_name,
            settings.EMBEDDING_MODEL,
        )

    @patch("apps.documents.services.SentenceTransformer")
    def test_rejects_embedding_dimension_mismatch(
        self,
        mock_sentence_transformer,
    ):
        model = Mock()

        model.get_embedding_dimension.return_value = (
            settings.EMBEDDING_DIMENSION + 1
        )

        mock_sentence_transformer.return_value = model

        with self.assertRaises(
            EmbeddingConfigurationError,
        ) as context:
            EmbeddingService()

        self.assertEqual(
            context.exception.code,
            "embedding_dimension_mismatch",
        )

    @patch("apps.documents.services.SentenceTransformer")
    def test_handles_model_load_failure(
        self,
        mock_sentence_transformer,
    ):
        mock_sentence_transformer.side_effect = RuntimeError(
            "model unavailable",
        )

        with self.assertRaises(
            EmbeddingConfigurationError,
        ) as context:
            EmbeddingService()

        self.assertEqual(
            context.exception.code,
            "embedding_model_load_failed",
        )

    @patch("apps.documents.services.SentenceTransformer")
    def test_embeds_multiple_documents(
        self,
        mock_sentence_transformer,
    ):
        model = Mock()

        model.get_embedding_dimension.return_value = (
            settings.EMBEDDING_DIMENSION
        )

        vectors = np.zeros(
            (
                2,
                settings.EMBEDDING_DIMENSION,
            ),
            dtype=np.float32,
        )

        vectors[0][0] = 1.0
        vectors[1][1] = 1.0

        model.encode.return_value = vectors

        mock_sentence_transformer.return_value = model

        service = EmbeddingService()

        result = service.embed_documents(
            [
                "متن اول",
                "متن دوم",
            ]
        )

        self.assertEqual(
            len(result),
            2,
        )

        self.assertEqual(
            len(result[0]),
            settings.EMBEDDING_DIMENSION,
        )

        self.assertEqual(
            len(result[1]),
            settings.EMBEDDING_DIMENSION,
        )

        model.encode.assert_called_once()

        encode_kwargs = (
            model.encode.call_args.kwargs
        )

        self.assertTrue(
            encode_kwargs["normalize_embeddings"],
        )

        self.assertTrue(
            encode_kwargs["convert_to_numpy"],
        )

    @patch("apps.documents.services.SentenceTransformer")
    def test_rejects_invalid_embedding_shape(
        self,
        mock_sentence_transformer,
    ):
        model = Mock()

        model.get_embedding_dimension.return_value = (
            settings.EMBEDDING_DIMENSION
        )

        model.encode.return_value = np.zeros(
            settings.EMBEDDING_DIMENSION,
            dtype=np.float32,
        )

        mock_sentence_transformer.return_value = model

        service = EmbeddingService()

        with self.assertRaises(
            EmbeddingConfigurationError,
        ) as context:
            service.embed_documents(
                ["text"],
            )

        self.assertEqual(
            context.exception.code,
            "invalid_embedding_shape",
        )

    @patch("apps.documents.services.SentenceTransformer")
    def test_empty_input_returns_empty_output(
        self,
        mock_sentence_transformer,
    ):
        model = Mock()

        model.get_embedding_dimension.return_value = (
            settings.EMBEDDING_DIMENSION
        )

        mock_sentence_transformer.return_value = model

        service = EmbeddingService()

        result = service.embed_documents([])

        self.assertEqual(
            result,
            [],
        )

        model.encode.assert_not_called()

    @patch("apps.documents.services.EmbeddingService")
    def test_embedding_service_is_singleton(
        self,
        mock_embedding_service,
    ):
        get_embedding_service.cache_clear()

        first = Mock()
        second = first

        mock_embedding_service.return_value = first

        service_one = get_embedding_service()
        service_two = get_embedding_service()

        self.assertIs(
            service_one,
            service_two,
        )

        mock_embedding_service.assert_called_once()


class DocumentServiceTests(
    DocumentTestMixin,
    TemporaryMediaMixin,
    TestCase,
):
    def test_create_document_creates_complete_state(self):
        with self.mocked_embedding_service():
            content = self.make_docx(
                paragraphs=[
                    "متن اول " * 100,
                    "متن دوم " * 100,
                ],
            )

            document = create_document(
                title="Test Document",
                uploaded_file=self.make_upload(
                    content,
                ),
            )

        self.assertEqual(
            document.title,
            "Test Document",
        )

        self.assertTrue(
            document.full_text,
        )

        self.assertGreater(
            document.chunk_count,
            0,
        )

        self.assertEqual(
            document.chunks.count(),
            document.chunk_count,
        )

        self.assertTrue(
            document.file.name.startswith(
                "documents/",
            )
        )

        self.assertTrue(
            document.file.name.endswith(
                ".docx",
            )
        )

        first_chunk = (
            document.chunks.order_by(
                "chunk_index",
            ).first()
        )

        self.assertIsNotNone(
            first_chunk,
        )

        self.assertEqual(
            len(first_chunk.embedding),
            settings.EMBEDDING_DIMENSION,
        )

        self.assertEqual(
            document.embedding_model,
            settings.EMBEDDING_MODEL,
        )

        self.assertTrue(
            default_storage.exists(
                document.file.name,
            )
        )

    def test_create_document_stores_vector_in_postgres(self):
        with self.mocked_embedding_service():
            content = self.make_docx(
                paragraphs=[
                    "متن برای تست pgvector " * 150,
                ],
            )

            document = create_document(
                title="Vector Test",
                uploaded_file=self.make_upload(
                    content,
                ),
            )

        chunk = (
            document.chunks
            .order_by("chunk_index")
            .first()
        )

        with connection.cursor() as cursor:
            cursor.execute(
                """
                SELECT vector_dims(embedding)
                FROM documents_documentchunk
                WHERE id = %s
                """,
                [chunk.id],
            )

            dimension = cursor.fetchone()[0]

        self.assertEqual(
            dimension,
            settings.EMBEDDING_DIMENSION,
        )

    def test_create_document_rolls_back_database_and_file(
        self,
    ):
        # Other tests in this class may leave files on disk
        # (TestCase rolls back DB rows, not filesystem writes).
        files_before = {
            path
            for path in Path(settings.MEDIA_ROOT).rglob("*")
            if path.is_file()
        }

        with self.mocked_embedding_service():
            content = self.make_docx(
                paragraphs=[
                    "متن rollback " * 100,
                ],
            )

            with patch(
                "apps.documents.services._replace_chunks",
                side_effect=RuntimeError(
                    "forced persistence failure",
                ),
            ):
                with self.assertRaises(
                    RuntimeError,
                ):
                    create_document(
                        title="Rollback Test",
                        uploaded_file=self.make_upload(
                            content,
                        ),
                    )

        self.assertEqual(
            Document.objects.count(),
            0,
        )

        self.assertEqual(
            DocumentChunk.objects.count(),
            0,
        )

        files_after = {
            path
            for path in Path(settings.MEDIA_ROOT).rglob("*")
            if path.is_file()
        }

        # The failed create must not leave a new orphaned file.
        self.assertEqual(
            files_after,
            files_before,
        )

    def test_update_title_does_not_reindex(self):
        with self.mocked_embedding_service() as embedding_service:
            content = self.make_docx(
                paragraphs=[
                    "متن اصلی " * 100,
                ],
            )

            document = create_document(
                title="Original",
                uploaded_file=self.make_upload(
                    content,
                ),
            )

            embedding_service.embed_documents.reset_mock()

            updated = update_document(
                document,
                title="Updated",
            )

        self.assertEqual(
            updated.title,
            "Updated",
        )

        self.assertEqual(
            embedding_service.embed_documents.call_count,
            0,
        )

        self.assertGreater(
            updated.chunks.count(),
            0,
        )

    def test_document_replacement_removes_old_chunks_and_file(
        self,
    ):
        with self.mocked_embedding_service():
            old_content = self.make_docx(
                paragraphs=[
                    "متن اولیه " * 100,
                ],
            )

            document = create_document(
                title="Original",
                uploaded_file=self.make_upload(
                    old_content,
                ),
            )

            old_file_name = document.file.name
            old_chunk_ids = set(
                document.chunks.values_list(
                    "id",
                    flat=True,
                )
            )

            new_content = self.make_docx(
                paragraphs=[
                    "متن جدید " * 200,
                ],
            )

            updated = update_document(
                document,
                title="Updated",
                uploaded_file=self.make_upload(
                    new_content,
                ),
            )

        new_chunk_ids = set(
            updated.chunks.values_list(
                "id",
                flat=True,
            )
        )

        self.assertEqual(
            updated.title,
            "Updated",
        )

        self.assertTrue(
            updated.full_text.startswith(
                "متن جدید",
            )
        )

        self.assertTrue(
            new_chunk_ids,
        )

        self.assertTrue(
            old_chunk_ids.isdisjoint(
                new_chunk_ids,
            )
        )

        self.assertEqual(
            updated.chunk_count,
            updated.chunks.count(),
        )

        self.assertNotEqual(
            old_file_name,
            updated.file.name,
        )

        self.assertFalse(
            default_storage.exists(
                old_file_name,
            )
        )

        self.assertTrue(
            default_storage.exists(
                updated.file.name,
            )
        )

    def test_reindex_rebuilds_chunks_from_existing_file(
        self,
    ):
        with self.mocked_embedding_service():
            content = self.make_docx(
                paragraphs=[
                    "متن قابل reindex " * 200,
                ],
            )

            document = create_document(
                title="Reindex Test",
                uploaded_file=self.make_upload(
                    content,
                ),
            )

            original_file_name = document.file.name
            original_chunk_ids = set(
                document.chunks.values_list(
                    "id",
                    flat=True,
                )
            )

            reindexed = reindex_document(
                document,
            )

        new_chunk_ids = set(
            reindexed.chunks.values_list(
                "id",
                flat=True,
            )
        )

        self.assertEqual(
            reindexed.file.name,
            original_file_name,
        )

        self.assertTrue(
            original_chunk_ids.isdisjoint(
                new_chunk_ids,
            )
        )

        self.assertEqual(
            reindexed.chunk_count,
            reindexed.chunks.count(),
        )

        self.assertTrue(
            reindexed.full_text,
        )

        self.assertTrue(
            default_storage.exists(
                reindexed.file.name,
            )
        )

    def test_reindex_reports_missing_source_file(
        self,
    ):
        with self.mocked_embedding_service():
            content = self.make_docx(
                paragraphs=[
                    "متن فایل " * 100,
                ],
            )

            document = create_document(
                title="Missing Source",
                uploaded_file=self.make_upload(
                    content,
                ),
            )

            old_chunk_ids = set(
                document.chunks.values_list(
                    "id",
                    flat=True,
                )
            )

            default_storage.delete(
                document.file.name,
            )

            with self.assertRaises(
                DocumentIngestionError,
            ) as context:
                reindex_document(document)

        self.assertEqual(
            context.exception.code,
            "source_file_missing",
        )

        self.assertEqual(
            set(
                DocumentChunk.objects.filter(
                    document=document,
                ).values_list(
                    "id",
                    flat=True,
                )
            ),
            old_chunk_ids,
        )

    def test_delete_document_removes_database_rows_and_file(
        self,
    ):
        with self.mocked_embedding_service():
            content = self.make_docx(
                paragraphs=[
                    "متن برای حذف " * 100,
                ],
            )

            document = create_document(
                title="Delete Test",
                uploaded_file=self.make_upload(
                    content,
                ),
            )

        file_name = document.file.name

        self.assertTrue(
            default_storage.exists(
                file_name,
            )
        )

        delete_document(document)

        self.assertFalse(
            Document.objects.filter(
                pk=document.pk,
            ).exists()
        )

        self.assertFalse(
            DocumentChunk.objects.filter(
                document_id=document.pk,
            ).exists()
        )

        self.assertFalse(
            default_storage.exists(
                file_name,
            )
        )


class DocumentAPITests(
    DocumentTestMixin,
    TemporaryMediaMixin,
    APITestCase,
):
    def setUp(self):
        self.client = APIClient()

    def test_create_document_api(self):
        with self.mocked_embedding_service():
            content = self.make_docx(
                paragraphs=[
                    "محتوای سند آزمایشی " * 100,
                ],
            )

            response = self.client.post(
                "/api/documents/",
                {
                    "title": "سند آزمایشی",
                    "file": self.make_upload(content),
                },
                format="multipart",
            )

        self.assertEqual(
            response.status_code,
            status.HTTP_201_CREATED,
        )

        document = Document.objects.get(
            pk=response.data["id"],
        )

        self.assertEqual(
            document.chunk_count,
            document.chunks.count(),
        )

        self.assertEqual(
            response.data["chunk_count"],
            document.chunk_count,
        )

    def test_create_requires_file(self):
        response = self.client.post(
            "/api/documents/",
            {
                "title": "بدون فایل",
            },
            format="multipart",
        )

        self.assertEqual(
            response.status_code,
            status.HTTP_400_BAD_REQUEST,
        )

        self.assertEqual(
            response.data["error"]["code"],
            "validation_error",
        )

    def test_reject_invalid_file(self):
        response = self.client.post(
            "/api/documents/",
            {
                "title": "Invalid",
                "file": self.make_upload(
                    b"not a docx",
                    name="evil.exe",
                ),
            },
            format="multipart",
        )

        self.assertEqual(
            response.status_code,
            status.HTTP_400_BAD_REQUEST,
        )

        self.assertEqual(
            response.data["error"]["code"],
            "validation_error",
        )

    def test_list_documents(self):
        with self.mocked_embedding_service():
            content = self.make_docx(
                paragraphs=[
                    "hello " * 100,
                ],
            )

            self.client.post(
                "/api/documents/",
                {
                    "title": "Document One",
                    "file": self.make_upload(content),
                },
                format="multipart",
            )

        response = self.client.get(
            "/api/documents/",
        )

        self.assertEqual(
            response.status_code,
            status.HTTP_200_OK,
        )

        self.assertEqual(
            len(response.data["results"]),
            1,
        )

    def test_retrieve_document(self):
        with self.mocked_embedding_service():
            content = self.make_docx(
                paragraphs=[
                    "متن برای retrieve " * 100,
                ],
            )

            create_response = self.client.post(
                "/api/documents/",
                {
                    "title": "Retrieve Test",
                    "file": self.make_upload(content),
                },
                format="multipart",
            )

        document_id = create_response.data["id"]

        response = self.client.get(
            f"/api/documents/{document_id}/",
        )

        self.assertEqual(
            response.status_code,
            status.HTTP_200_OK,
        )

        self.assertEqual(
            response.data["id"],
            document_id,
        )

        self.assertIn(
            "متن برای retrieve",
            response.data["full_text"],
        )

    def test_patch_title(self):
        with self.mocked_embedding_service():
            content = self.make_docx(
                paragraphs=[
                    "متن patch " * 100,
                ],
            )

            create_response = self.client.post(
                "/api/documents/",
                {
                    "title": "Original",
                    "file": self.make_upload(content),
                },
                format="multipart",
            )

        document_id = create_response.data["id"]

        response = self.client.patch(
            f"/api/documents/{document_id}/",
            {
                "title": "Updated",
            },
            format="multipart",
        )

        self.assertEqual(
            response.status_code,
            status.HTTP_200_OK,
        )

        self.assertEqual(
            response.data["title"],
            "Updated",
        )

    def test_patch_file_reindexes_document(self):
        with self.mocked_embedding_service():
            old_content = self.make_docx(
                paragraphs=[
                    "متن قدیمی " * 100,
                ],
            )

            create_response = self.client.post(
                "/api/documents/",
                {
                    "title": "Original",
                    "file": self.make_upload(
                        old_content,
                    ),
                },
                format="multipart",
            )

            document_id = create_response.data["id"]

            new_content = self.make_docx(
                paragraphs=[
                    "متن جدید " * 150,
                ],
            )

            response = self.client.patch(
                f"/api/documents/{document_id}/",
                {
                    "title": "Updated",
                    "file": self.make_upload(
                        new_content,
                    ),
                },
                format="multipart",
            )

        self.assertEqual(
            response.status_code,
            status.HTTP_200_OK,
        )

        self.assertEqual(
            response.data["title"],
            "Updated",
        )

        self.assertIn(
            "متن جدید",
            response.data["full_text"],
        )

        self.assertNotIn(
            "متن قدیمی",
            response.data["full_text"],
        )

    def test_reindex_endpoint(self):
        with self.mocked_embedding_service():
            content = self.make_docx(
                paragraphs=[
                    "متن reindex endpoint " * 100,
                ],
            )

            create_response = self.client.post(
                "/api/documents/",
                {
                    "title": "Reindex API",
                    "file": self.make_upload(content),
                },
                format="multipart",
            )

            document_id = create_response.data["id"]

            response = self.client.post(
                f"/api/documents/{document_id}/reindex/",
            )

        self.assertEqual(
            response.status_code,
            status.HTTP_200_OK,
        )

        self.assertGreater(
            response.data["chunk_count"],
            0,
        )

    def test_delete_endpoint(self):
        with self.mocked_embedding_service():
            content = self.make_docx(
                paragraphs=[
                    "متن delete endpoint " * 100,
                ],
            )

            create_response = self.client.post(
                "/api/documents/",
                {
                    "title": "Delete API",
                    "file": self.make_upload(content),
                },
                format="multipart",
            )

        document_id = create_response.data["id"]

        response = self.client.delete(
            f"/api/documents/{document_id}/",
        )

        self.assertEqual(
            response.status_code,
            status.HTTP_204_NO_CONTENT,
        )

        self.assertFalse(
            Document.objects.filter(
                pk=document_id,
            ).exists()
        )

    def test_document_not_found(self):
        response = self.client.get(
            "/api/documents/999999/",
        )

        self.assertEqual(
            response.status_code,
            status.HTTP_404_NOT_FOUND,
        )

        self.assertEqual(
            response.data["error"]["code"],
            "not_found",
        )


class RealEmbeddingIntegrationTests(
    DocumentTestMixin,
    TemporaryMediaMixin,
    TestCase,
):
    """
    Verify the complete ingestion pipeline with the real
    SentenceTransformer model and PostgreSQL/pgvector.
    """

    @classmethod
    def setUpClass(cls):
        super().setUpClass()

        get_embedding_service.cache_clear()

    @classmethod
    def tearDownClass(cls):
        get_embedding_service.cache_clear()

        super().tearDownClass()

    def test_real_docx_to_embedding_to_pgvector(self):
        content = self.make_docx(
            paragraphs=[
                (
                    "این یک سند واقعی برای تست "
                    "سیستم Document Q&A است. "
                )
                * 150,
                (
                    "این بخش برای بررسی embedding "
                    "متن فارسی استفاده می‌شود. "
                )
                * 150,
            ],
            table_rows=[
                ["موضوع", "مقدار"],
                ["Embedding", "384"],
                ["Vector DB", "PostgreSQL + pgvector"],
            ],
        )

        uploaded_file = self.make_upload(
            content,
            name="real-integration-test.docx",
        )

        document = create_document(
            title="Real Embedding Integration Test",
            uploaded_file=uploaded_file,
        )

        self.assertTrue(
            document.full_text,
        )

        self.assertIn(
            "سیستم Document Q&A",
            document.full_text,
        )

        self.assertGreater(
            document.chunk_count,
            0,
        )

        self.assertEqual(
            document.chunk_count,
            document.chunks.count(),
        )

        chunks = list(
            document.chunks.order_by(
                "chunk_index",
            )
        )

        self.assertEqual(
            len(chunks),
            document.chunk_count,
        )

        for chunk in chunks:
            self.assertEqual(
                len(chunk.embedding),
                settings.EMBEDDING_DIMENSION,
            )

        first_vector = chunks[0].embedding

        norm = math.sqrt(
            sum(
                value * value
                for value in first_vector
            )
        )

        self.assertAlmostEqual(
            norm,
            1.0,
            places=3,
        )

        with connection.cursor() as cursor:
            cursor.execute(
                """
                SELECT
                    vector_dims(embedding),
                    COUNT(*)
                FROM documents_documentchunk
                WHERE document_id = %s
                GROUP BY vector_dims(embedding)
                """,
                [document.id],
            )

            result = cursor.fetchone()

        self.assertIsNotNone(
            result,
        )

        vector_dimension, chunk_count = result

        self.assertEqual(
            vector_dimension,
            settings.EMBEDDING_DIMENSION,
        )

        self.assertEqual(
            chunk_count,
            document.chunk_count,
        )

        self.assertEqual(
            document.embedding_model,
            settings.EMBEDDING_MODEL,
        )

        self.assertTrue(
            default_storage.exists(
                document.file.name,
            )
        )



class TokenizerTests(TestCase):
    def test_persian_tokenization(self):
        tokens = tokenize_text("شرایط مرخصی کارکنان چیست؟")
        self.assertIn("شرایط", tokens)
        self.assertIn("مرخصی", tokens)
        self.assertIn("کارکنان", tokens)

    def test_english_tokenization(self):
        tokens = tokenize_text("What is the annual leave policy?")
        self.assertIn("what", tokens)
        self.assertIn("annual", tokens)
        self.assertIn("leave", tokens)

    def test_mixed_tokenization(self):
        tokens = tokenize_text("مدت leave سالانه چقدر است؟")
        self.assertIn("مدت", tokens)
        self.assertIn("leave", tokens)
        self.assertIn("سالانه", tokens)

    def test_punctuation_stripped(self):
        tokens = tokenize_text("hello, world!!!")
        self.assertEqual(tokens, ["hello", "world"])

    def test_persian_yeh_kaf_normalization(self):
        tokens = tokenize_text("يك كتاب")
        self.assertIn("یک", tokens)
        self.assertIn("کتاب", tokens)


class ScoringTests(TestCase):
    def test_vector_score_formula(self):
        self.assertAlmostEqual(
            compute_vector_score(0.0),
            1.0,
        )
        self.assertAlmostEqual(
            compute_vector_score(1.0),
            0.5,
        )

    def test_keyword_score_overlap(self):
        question_terms = tokenize_text(
            "مرخصی سالانه کارکنان",
        )
        score = compute_keyword_score(
            question_terms,
            "مرخصی کارکنان سالانه بیست روز است.",
        )
        self.assertGreater(score, 0.5)

    def test_keyword_score_unrelated(self):
        question_terms = tokenize_text(
            "مرخصی سالانه",
        )
        score = compute_keyword_score(
            question_terms,
            "آب و هوا امروز آفتابی است.",
        )
        self.assertEqual(score, 0.0)

    def test_final_score_weights(self):
        final = compute_final_score(1.0, 0.0)
        self.assertAlmostEqual(final, 0.7)

        final = compute_final_score(0.0, 1.0)
        self.assertAlmostEqual(final, 0.3)


class RetrievalValidationTests(TestCase):
    def test_empty_question_rejected(self):
        with self.assertRaises(RetrievalError) as ctx:
            retrieve_chunks(question="")
        self.assertEqual(ctx.exception.code, "empty_question")

    def test_whitespace_question_rejected(self):
        with self.assertRaises(RetrievalError) as ctx:
            retrieve_chunks(question="   ")
        self.assertEqual(ctx.exception.code, "empty_question")

    def test_top_k_zero_rejected(self):
        with self.assertRaises(RetrievalError) as ctx:
            retrieve_chunks(question="سوال", top_k=0)
        self.assertEqual(ctx.exception.code, "invalid_top_k")

    def test_top_k_negative_rejected(self):
        with self.assertRaises(RetrievalError) as ctx:
            retrieve_chunks(question="سوال", top_k=-1)
        self.assertEqual(ctx.exception.code, "invalid_top_k")

    def test_top_k_too_large_rejected(self):
        with self.assertRaises(RetrievalError) as ctx:
            retrieve_chunks(question="سوال", top_k=1000)
        self.assertEqual(ctx.exception.code, "invalid_top_k")

    def test_empty_document_ids_rejected(self):
        with self.assertRaises(RetrievalError) as ctx:
            retrieve_chunks(
                question="سوال",
                document_ids=[],
            )
        self.assertEqual(
            ctx.exception.code,
            "empty_document_ids",
        )


class RetrievalServiceTests(
    DocumentTestMixin,
    TemporaryMediaMixin,
    TestCase,
):
    def _seed_chunks(self):
        """
        Create three documents with known text via mocked embeddings.
        """
        with self.mocked_embedding_service(value=0.01):
            docs = []
            for title, paragraph in [
                (
                    "Leave Policy",
                    "مرخصی کارکنان سالانه بیست روز است.",
                ),
                (
                    "Work Hours",
                    "ساعات کاری شرکت از ساعت هشت صبح تا چهار بعدازظهر است.",
                ),
                (
                    "Contract",
                    "قرارداد کارکنان به صورت یک ساله تنظیم می‌شود.",
                ),
            ]:
                content = self.make_docx(
                    paragraphs=[paragraph * 20],
                )
                document = create_document(
                    title=title,
                    uploaded_file=self.make_upload(content),
                )
                docs.append(document)
            return docs

    def test_retrieve_ranks_relevant_document_first(self):
        docs = self._seed_chunks()

        # Fixed embedding so ranking relies on hybrid + text content
        # Use distinct embeddings per call would need more complex mock.
        # Instead patch CosineDistance path by giving all same vector score
        # and verify keyword pushes leave policy up.
        with self.mocked_embedding_service(value=0.02):
            results = retrieve_chunks(
                question="میزان مرخصی سالانه کارکنان چقدر است؟",
                top_k=3,
            )

        self.assertTrue(results)
        self.assertEqual(
            results[0].document_title,
            "Leave Policy",
        )
        self.assertEqual(
            results[0].document_id,
            docs[0].id,
        )
        self.assertLessEqual(len(results), 3)

    def test_document_ids_filter(self):
        docs = self._seed_chunks()

        with self.mocked_embedding_service(value=0.02):
            results = retrieve_chunks(
                question="ساعات کاری شرکت چیست؟",
                top_k=5,
                document_ids=[docs[1].id],
            )

        self.assertTrue(results)
        for item in results:
            self.assertEqual(item.document_id, docs[1].id)

    def test_top_k_limits_results(self):
        self._seed_chunks()

        with self.mocked_embedding_service(value=0.02):
            results = retrieve_chunks(
                question="کارکنان شرکت",
                top_k=1,
            )

        self.assertEqual(len(results), 1)

    def test_result_metadata_present(self):
        docs = self._seed_chunks()

        with self.mocked_embedding_service(value=0.02):
            results = retrieve_chunks(
                question="مرخصی کارکنان",
                top_k=3,
            )

        self.assertTrue(results)
        item = results[0]
        self.assertIsInstance(item.document_id, int)
        self.assertTrue(item.document_title)
        self.assertIsInstance(item.chunk_index, int)
        self.assertTrue(item.text)
        self.assertIsInstance(item.distance, float)
        self.assertIsInstance(item.vector_score, float)
        self.assertIsInstance(item.keyword_score, float)
        self.assertIsInstance(item.final_score, float)

        # final_score formula check roughly
        expected = 0.7 * item.vector_score + 0.3 * item.keyword_score
        self.assertAlmostEqual(item.final_score, expected, places=5)

    def test_empty_database_returns_empty(self):
        with self.mocked_embedding_service(value=0.02):
            results = retrieve_chunks(
                question="هر سوالی",
                top_k=6,
            )
        self.assertEqual(results, [])

    def test_results_sorted_by_final_score(self):
        self._seed_chunks()

        with self.mocked_embedding_service(value=0.02):
            results = retrieve_chunks(
                question="کارکنان",
                top_k=6,
            )

        scores = [item.final_score for item in results]
        self.assertEqual(scores, sorted(scores, reverse=True))

    def test_uses_embedding_service(self):
        self._seed_chunks()

        with self.mocked_embedding_service(value=0.05) as mock_svc:
            retrieve_chunks(
                question="تست embedding",
                top_k=2,
            )
            mock_svc.embed_documents.assert_called()


class RetrievalAPITests(
    DocumentTestMixin,
    TemporaryMediaMixin,
    APITestCase,
):
    def setUp(self):
        self.client = APIClient()

    def test_retrieval_endpoint(self):
        with self.mocked_embedding_service(value=0.02):
            content = self.make_docx(
                paragraphs=[
                    "مرخصی کارکنان سالانه بیست روز است. " * 30,
                ],
            )
            self.client.post(
                "/api/documents/",
                {
                    "title": "Policy",
                    "file": self.make_upload(content),
                },
                format="multipart",
            )

            response = self.client.post(
                "/api/retrieval/",
                {
                    "question": "مرخصی سالانه چقدر است؟",
                    "top_k": 3,
                },
                format="json",
            )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertIn("results", response.data)
        self.assertIn("question", response.data)
        if response.data["results"]:
            item = response.data["results"][0]
            self.assertIn("document_id", item)
            self.assertIn("final_score", item)
            self.assertIn("text", item)

    def test_retrieval_rejects_empty_question(self):
        response = self.client.post(
            "/api/retrieval/",
            {"question": "   "},
            format="json",
        )
        self.assertEqual(
            response.status_code,
            status.HTTP_400_BAD_REQUEST,
        )

